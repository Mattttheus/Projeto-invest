"""Aba Meta: renda mensal que cada ação e cada FII deve pagar (META_MENSAL, aba Premissas),
quantas cotas são necessárias e quanto custa chegar lá — total e o que ainda falta.

Tudo por fórmula, lendo a Carteira (mesma linha de cada ativo): mudar a meta, o aporte ou as cotações
recalcula a aba inteira.
"""
from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import get_column_letter as L

from .abas import D0, H
from .estilo import (BRL, BRL0, BRL4, MESES, PCT, QTD, TEAL, cabecalho, cel, f_bold, f_link, f_norm, faixa,
                     fill_total, fill_zebra, larguras, linha_total, notas, tabela)

# (cabeçalho, fórmula da linha r, formato, vem da Carteira?)
COLUNAS = [
    ("Ticker", '=Carteira!A{r}', None, True),
    ("Tipo", '=Carteira!B{r}', None, True),
    ("Segmento", '=Carteira!AB{r}', None, True),
    ("Cotação (R$)", '=Carteira!F{r}', BRL, True),
    ("Provento 12m por cota (R$)", '=Carteira!T{r}', BRL4, True),
    ("Provento mensal por cota (R$)", '=Carteira!U{r}', BRL4, True),
    ("Rendimento mensal", '=IF(D{r}>0,F{r}/D{r},0)', '0.00%', False),
    ("Meta mensal (R$)", '=IF(A{r}="","",META_MENSAL)', BRL, False),
    ("Cotas necessárias", '=IF(A{r}="","",IF(F{r}>0,ROUNDUP(H{r}/F{r},0),"sem proventos"))', QTD, False),
    ("Custo total da meta (R$)", '=IF(ISNUMBER(I{r}),I{r}*D{r},"")', BRL0, False),
    ("Cotas que você tem", '=Carteira!C{r}', QTD, True),
    ("Renda mensal atual (R$)", '=Carteira!V{r}', BRL, True),
    ("% da meta", '=IF(N(H{r})>0,MIN(1,L{r}/H{r}),0)', PCT, False),
    ("Cotas faltantes", '=IF(ISNUMBER(I{r}),MAX(0,I{r}-K{r}),"")', QTD, False),
    ("Custo faltante (R$)", '=IF(ISNUMBER(N{r}),N{r}*D{r},"")', BRL0, False),
    ("Meses (aporte só neste ativo)", '=IF(AND(ISNUMBER(O{r}),APORTE_MENSAL>0),ROUNDUP(O{r}/APORTE_MENSAL,0),"")', MESES, False),
]
SOMAS = ("H", "J", "L", "N", "O")                    # colunas somadas no total e nos subtotais


def aba_meta(wb, N, T):
    ws = wb.create_sheet("Meta")
    n = len(COLUNAS)
    faixa(ws, "Meta de Renda por Ativo", "Cada ação e cada FII deve pagar a meta mensal da aba Premissas — cotas e custo "
          "para chegar lá. Calculado, não digite aqui.", n)
    cabecalho(ws, H, [c[0] for c in COLUNAS])
    ws.row_dimensions[H].height = 44
    for i in range(N):
        r = D0 + i
        z = fill_zebra if i % 2 else None
        for col, (_, formula, fmt, vinculo) in enumerate(COLUNAS, 1):
            cel(ws, r, col, formula.format(r=r), fmt, f_link if vinculo else f_norm, z)
    tabela(ws, "tMeta", H, T - 1, n)
    _totais(ws, T)
    _subtotais(ws, T)
    ws.conditional_formatting.add(f"M{D0}:M{T - 1}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                                 end_value=1, color=TEAL))
    notas(ws, T + 6, [
        "Cotas necessárias = meta mensal ÷ provento mensal por cota (provento dos últimos 12 meses ÷ 12), arredondado para cima.",
        "Custo total = cotas necessárias × cotação de hoje; custo faltante considera as cotas que você já tem.",
        "Meses: tempo até a meta se todo o aporte mensal (Premissas) fosse para este ativo. O total considera a carteira inteira.",
        "Ativo sem proventos nos últimos 12 meses não tem como atingir a meta de renda ('sem proventos').",
        "Para mudar a meta, edite meta_mensal_por_ativo em dados/config.json (ou a aba Premissas).",
    ])
    larguras(ws, [10, 7, 22, 12, 14, 14, 13, 12, 12, 16, 12, 14, 10, 12, 16, 14])
    ws.freeze_panes = "B5"
    return ws


def _totais(ws, T):
    linha_total(ws, T, len(COLUNAS))
    for c in SOMAS:
        fmt = QTD if c == "N" else (BRL if c in ("H", "L") else BRL0)
        ws[f"{c}{T}"] = f"=SUM({c}{D0}:{c}{T - 1})"
        ws[f"{c}{T}"].number_format = fmt
    ws[f"M{T}"] = f"=IF(H{T}>0,L{T}/H{T},0)"
    ws[f"M{T}"].number_format = PCT
    ws[f"P{T}"] = f'=IF(APORTE_MENSAL>0,ROUNDUP(O{T}/APORTE_MENSAL,0),"")'
    ws[f"P{T}"].number_format = MESES


def _subtotais(ws, T):
    for k, (tipo, rotulo) in enumerate((("AÇÃO", "Subtotal ações"), ("FII", "Subtotal FIIs"))):
        r = T + 2 + k
        ws.cell(r, 1, rotulo).font = f_bold
        for c in SOMAS:
            fmt = QTD if c == "N" else BRL0
            cel(ws, r, _num(c), f'=SUMIFS({c}{D0}:{c}{T - 1},$B${D0}:$B${T - 1},"{tipo}")', fmt, f_bold, fill_total)
        cel(ws, r, _num("M"), f"=IF(H{r}>0,L{r}/H{r},0)", PCT, f_bold, fill_total)
        cel(ws, r, _num("P"), f'=IF(APORTE_MENSAL>0,ROUNDUP(O{r}/APORTE_MENSAL,0),"")', MESES, f_bold, fill_total)


def _num(letra):
    return next(i for i in range(1, len(COLUNAS) + 1) if L(i) == letra)
