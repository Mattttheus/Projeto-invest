"""Aba Painel no padrão "Casa Organizada" (tema claro), no Excel e no Google Planilhas.

Layout: faixa azul-marinho (marca, situação dos dados, botão dourado "Novo lançamento") | menu de abas
branco com a aba ativa sublinhada | painel "Filtros" à esquerda (tipo de ativo, atalhos, meta de renda)
| cards brancos de indicadores | gráficos brancos em azul/laranja | meta de renda | destaques | alocação.
Só lê as outras abas; o filtro Tipo (Todos / AÇÃO / FII) age sobre os rankings via Carteira!AC.
"""
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

from .estilo import BRL0, DOURADO, F, MESES, MODO, PCT, TOPO, topo

# ---- paleta clara (Casa Organizada): fundo cinza-claro, cards brancos, azul = dado principal, laranja = contraste
BG, CARD, BORDA_CARD, GRADE = "F1F5F9", "FFFFFF", "E2E8F0", "EEF2F6"
TXT, SUAVE, APAGADO = "0F172A", "64748B", "94A3B8"
AZUL_S, LARANJA_S, TEAL_S, CINZA_S = "2E6DB4", "D9772B", "0F9488", "94A3B8"
POSITIVO, NEGATIVO = "15803D", "DC2626"
ROSCA = [AZUL_S, LARANJA_S, TEAL_S, CINZA_S, TOPO, DOURADO]

# grade: A = margem | B..G = painel de filtros | H = respiro | I..AK = conteúdo (29 colunas) | AL = margem
ESQ0, ESQ1 = 2, 7
X0 = 9
ULT = X0 + 28
W3 = 9.3                                             # largura (cm) de cada gráfico (3 por linha, com respiro)
K = '"R$ "#,##0,"k";-"R$ "#,##0,"k";"R$ 0"'          # valores compactos: R$ 40k
F_BG, F_CARD = PatternFill("solid", fgColor=BG), PatternFill("solid", fgColor=CARD)
LADO = Side(style="thin", color=BORDA_CARD)

# posições usadas pelas macros (macros.py) para desenhar os botões
BOTAO_NOVO = f"{L(X0 + 19)}1:{L(X0 + 23)}1"
BOTAO_ATUALIZAR = f"{L(ULT - 4)}2:{L(ULT)}2"
FILTROS = {"Todos": ("B8:C8", "FiltrarTodos"), "AÇÃO": ("D8:E8", "FiltrarAcoes"), "FII": ("F8:G8", "FiltrarFIIs")}
BOTAO_LIMPAR = "F5:G5"
ATALHOS = [("✚  Lançar compra ou provento", "Cadastro"), ("◧  Carteira completa", "Carteira"),
           ("◎  Meta de renda por ativo", "Meta"),
           ("∑  Análise quantitativa", "Quant"), ("λ  Modelos matemáticos", "Modelos"),
           ("↗  Projeção futura", "Projeção"), ("⚙  Premissas e metas", "Premissas"),
           ("⚠  Verificar dados", "Verificar")]


# ------------------------------------------------------------------ gráficos claros

def _texto(tamanho=800, cor=SUAVE, negrito=False):
    cp = CharacterProperties(sz=tamanho, b=negrito, solidFill=cor)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _titulo(texto):
    cp = CharacterProperties(sz=1100, b=True, solidFill=TXT)
    rico = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=texto)])])
    return Title(tx=Text(rich=rico), overlay=False)


def _fundo(ch):
    ch.graphical_properties = GraphicalProperties(solidFill=CARD, ln=LineProperties(solidFill=BORDA_CARD))
    ch.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))


def barras(ws_dados, col, ws_cat, col_cat, r1, r2, titulo, fmt, cor, horizontal=True):
    """Barras/colunas de uma cor, com valor na ponta, em card branco."""
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.add_data(Reference(ws_dados, min_col=col, min_row=r1 - 1, max_row=r2), titles_from_data=True)
    ch.set_categories(Reference(ws_cat, min_col=col_cat, min_row=r1, max_row=r2))
    ch.y_axis.numFmt = ch.y_axis.number_format = fmt
    ch.gapWidth = 60
    ch.x_axis.tickLblSkip = 1
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"        # maior no topo
        ch.y_axis.crosses = "max"
        ch.x_axis.tickLblPos = "low"
    ch.legend = None
    ch.dataLabels = DataLabelList(showVal=True, showLegendKey=False, showCatName=False, showSerName=False,
                                  showPercent=False, dLblPos="outEnd")
    ch.dataLabels.txPr = _texto(800, "334155")
    ch.title = _titulo(titulo)
    ch.style = 2
    ch.width = W3
    _fundo(ch)
    s = ch.series[0]
    s.graphicalProperties.solidFill = cor
    s.graphicalProperties.line.solidFill = cor
    s.invertIfNegative = False
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.x_axis.txPr = ch.y_axis.txPr = _texto(800, SUAVE)
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRADE))
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=BORDA_CARD))
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    return ch


def rosca(ws_dados, col, ws_cat, col_cat, r1, r2, titulo):
    ch = DoughnutChart()
    ch.holeSize = 60
    ch.add_data(Reference(ws_dados, min_col=col, min_row=r1 - 1, max_row=r2), titles_from_data=True)
    ch.set_categories(Reference(ws_cat, min_col=col_cat, min_row=r1, max_row=r2))
    ch.dataLabels = DataLabelList(showPercent=True, showVal=False, showCatName=False, showSerName=False,
                                  showLegendKey=False, showLeaderLines=False)
    ch.dataLabels.txPr = _texto(800, "FFFFFF", True)
    ch.legend.position = "r"
    ch.legend.txPr = _texto(800, SUAVE)
    for i in range(r2 - r1 + 1):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = ROSCA[i % len(ROSCA)]
        pt.graphicalProperties.line.solidFill = CARD
        ch.series[0].dPt.append(pt)
    ch.title = _titulo(titulo)
    ch.style = 2
    ch.width = W3
    _fundo(ch)
    return ch


# ------------------------------------------------------------------ blocos da tela

def _pinta(ws, r1, r2, c1, c2, fill):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).fill = fill


def card(ws, r1, r2, c1, c2):
    """Retângulo branco com borda cinza-clara (base de todos os blocos)."""
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).fill = F_CARD
            ws.cell(r, c).border = Border(left=LADO if c == c1 else None, right=LADO if c == c2 else None,
                                          top=LADO if r == r1 else None, bottom=LADO if r == r2 else None)


def cabecalho(ws):
    """Faixa azul-marinho (marca + situação dos dados) e menu; linha 3 = data e fonte das cotações."""
    topo(ws, "Painel de investimentos", ULT + 1, altura=40)
    ws.merge_cells(start_row=1, start_column=X0 + 8, end_row=1, end_column=X0 + 17)
    st = ws.cell(1, X0 + 8, '=IF(COUNTIF(Verificar!A:A,"ERRO")>0,"⚠  Há erros nos dados — ver Verificar",'
                            'IF(COUNTIF(Verificar!A:A,"AVISO")>0,"●  "&COUNTIF(Verificar!A:A,"AVISO")&" pontos a conferir",'
                            '"✓  Dados consistentes"))')
    st.font = Font(name=F, size=9, bold=True, color="86EFAC")
    st.alignment = Alignment(vertical="center")
    st.hyperlink = "#Verificar!A1"
    for txt, cor in (("⚠", "FCA5A5"), ("●", "FCD34D")):
        ws.conditional_formatting.add(st.coordinate, FormulaRule(formula=[f'LEFT({st.coordinate},1)="{txt}"'],
                                                                 font=Font(name=F, bold=True, color=cor)))
    if MODO["google"]:                                   # sem macros: o botão vira um link
        ws.merge_cells(BOTAO_NOVO)
        b = ws[BOTAO_NOVO.split(":")[0]]
        b.value, b.hyperlink = "✚  Novo lançamento", "#Cadastro!A1"
        b.font = Font(name=F, size=10, bold=True, color=TOPO)
        b.fill = PatternFill("solid", fgColor=DOURADO)
        b.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 20
    ws.merge_cells(start_row=3, start_column=ESQ0, end_row=3, end_column=ULT)
    fonte = "Google Finance" if MODO["google"] else "Yahoo Finance"
    sub = ws.cell(3, ESQ0, "=TODAY()")                   # data pelo formato da célula (vale em qualquer idioma)
    sub.number_format = f'"Atualizado em "dd/mm/yyyy"   •   cotações: {fonte}   •   valores do filtro escolhido à esquerda"'
    sub.font = Font(name=F, size=8, color=SUAVE)
    sub.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[4].height = 8


def _rotulo(ws, r, texto, tamanho=11, cor=TXT, col=ESQ0):
    c = ws.cell(r, col, texto)
    c.font = Font(name=F, size=tamanho, bold=True, color=cor)
    c.alignment = Alignment(vertical="center", indent=1)
    return c


def filtros(ws, T, D0, fim):
    """Painel lateral: filtro Tipo (botões no Excel + lista), atalhos para os módulos e meta de renda."""
    card(ws, 5, fim, ESQ0, ESQ1)
    _rotulo(ws, 5, "Filtros")
    ws.row_dimensions[5].height = 24
    _rotulo(ws, 7, "Tipo de ativo", 9, SUAVE)
    if MODO["google"]:                                   # sem botões: os três tipos como legenda
        for tipo, (rng, _) in FILTROS.items():
            ws.merge_cells(rng)
            c = ws[rng.split(":")[0]]
            c.value = {"AÇÃO": "Ações", "FII": "FIIs"}.get(tipo, tipo)
            c.font = Font(name=F, size=9, color=SUAVE)
            c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[8].height = 24
    _rotulo(ws, 10, "Selecionado", 8, APAGADO)
    ws.merge_cells(start_row=10, start_column=ESQ0 + 2, end_row=10, end_column=ESQ1)
    filtro = ws.cell(10, ESQ0 + 2, "Todos")
    filtro.font = Font(name=F, size=10, bold=True, color=AZUL_S)
    filtro.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in range(ESQ0 + 2, ESQ1 + 1):
        ws.cell(10, c).fill = PatternFill("solid", fgColor="EFF4FB")
    dv = DataValidation(type="list", formula1='"Todos,AÇÃO,FII"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(filtro.coordinate)
    ws.parent.defined_names["FILTRO_TIPO"] = DefinedName("FILTRO_TIPO", attr_text=f"Painel!${L(ESQ0 + 2)}$10")
    ws.row_dimensions[10].height = 22

    _rotulo(ws, 12, "Atalhos")
    ws.row_dimensions[12].height = 24
    for k, (texto, aba) in enumerate(ATALHOS):
        r = 13 + k
        ws.merge_cells(start_row=r, start_column=ESQ0, end_row=r, end_column=ESQ1)
        c = ws.cell(r, ESQ0, texto)
        c.font = Font(name=F, size=9, bold=True, color="FFFFFF")
        c.alignment = Alignment(vertical="center", indent=1)
        c.hyperlink = f"#'{aba}'!A1" if not aba.isalnum() else f"#{aba}!A1"
        for cc in range(ESQ0, ESQ1 + 1):
            ws.cell(r, cc).fill = PatternFill("solid", fgColor=AZUL_S)
            ws.cell(r, cc).border = Border(bottom=Side(style="medium", color=CARD),
                                           left=LADO if cc == ESQ0 else None, right=LADO if cc == ESQ1 else None)
        ws.row_dimensions[r].height = 21

    r = 13 + len(ATALHOS) + 1
    _rotulo(ws, r, "Meta de renda")
    ws.row_dimensions[r].height = 24
    ws.merge_cells(start_row=r + 1, start_column=ESQ0, end_row=r + 1, end_column=ESQ1)
    p = ws.cell(r + 1, ESQ0, f"=Carteira!W{T}")
    p.number_format = '0%" da meta atingida"'
    p.font = Font(name=F, size=16, bold=True, color=TXT)
    p.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[r + 1].height = 28
    ws.merge_cells(start_row=r + 2, start_column=ESQ0, end_row=r + 2, end_column=ESQ1)
    barra = ws.cell(r + 2, ESQ0, f"=Carteira!W{T}")
    barra.number_format = ';;;'                          # só a barra, sem o número
    ws.conditional_formatting.add(barra.coordinate, DataBarRule(start_type="num", start_value=0, end_type="num",
                                                                end_value=1, color=DOURADO, showValue=False))
    ws.merge_cells(start_row=r + 3, start_column=ESQ0, end_row=r + 3, end_column=ESQ1)
    s = ws.cell(r + 3, ESQ0, f'="Renda de R$ "&FIXED(Carteira!V{T},0)&" por mês"')
    s.font = Font(name=F, size=8, color=SUAVE)
    s.alignment = Alignment(vertical="center", indent=1)


def kpi(ws, row, col, rotulo, valor, fmt, sub, sub_fmt=None, cor_sub=POSITIVO, span=4):
    """Card de indicador (como Renda / Despesas / Saldo): rótulo cinza, valor grande e linha de apoio colorida."""
    card(ws, row, row + 2, col, col + span - 1)
    for dr, (x, fnt, fm) in enumerate(((rotulo, Font(name=F, size=8, bold=True, color=SUAVE), None),
                                       (valor, Font(name=F, size=18, bold=True, color=TXT), fmt),
                                       (sub, Font(name=F, size=8, color=cor_sub), sub_fmt))):
        ws.merge_cells(start_row=row + dr, start_column=col, end_row=row + dr, end_column=col + span - 1)
        c = ws.cell(row + dr, col, x)
        c.font = fnt
        c.alignment = Alignment(vertical="center" if dr else "bottom", indent=1, shrink_to_fit=dr == 1)
        if fm:
            c.number_format = fm
    ws.row_dimensions[row].height = 20
    ws.row_dimensions[row + 1].height = 32
    ws.row_dimensions[row + 2].height = 18
    return ws.cell(row + 1, col)


def _altura_cm(ws, r1, r2):
    """Altura (cm) das linhas r1..r2 — o gráfico ocupa o bloco inteiro, sem vão embaixo."""
    pts = sum(ws.row_dimensions[r].height or 15 for r in range(r1, r2 + 1))
    return pts * 2.54 / 72 - 0.25


def _grafico(ws, ch, celula, r1, r2):
    ch.height = _altura_cm(ws, r1, r2)
    ws.add_chart(ch, celula)


def secao(ws, row, texto, dica=""):
    c = ws.cell(row, X0, texto)
    c.font = Font(name=F, size=11, bold=True, color=TXT)
    if dica:
        d = ws.cell(row, X0 + 12, dica)
        d.font = Font(name=F, size=8, color=SUAVE)
    ws.row_dimensions[row].height = 22


def sinal(ws, faixa, tamanho=None):
    for op, cor in (("lessThan", NEGATIVO), ("greaterThan", POSITIVO)):
        ws.conditional_formatting.add(faixa, CellIsRule(operator=op, formula=["0"],
                                                        font=Font(name=F, size=tamanho, bold=bool(tamanho), color=cor)))


def destaque(ws, row, col, titulo, cabecalhos, linhas):
    """Card com mini-tabela de 9 colunas: #, ticker e dois valores."""
    spans = [(col, col), (col + 1, col + 2), (col + 3, col + 5), (col + 6, col + 8)]
    card(ws, row, row + 1 + len(linhas), col, col + 8)
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 8)
    t = ws.cell(row, col, titulo)
    t.font = Font(name=F, size=10, bold=True, color=TXT)
    t.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[row].height = 24
    for (a, b), h in zip(spans, ["#", "Ativo", *cabecalhos]):
        ws.merge_cells(start_row=row + 1, start_column=a, end_row=row + 1, end_column=b)
        cc = ws.cell(row + 1, a, h.upper())
        cc.font = Font(name=F, size=7, bold=True, color=SUAVE)
        cc.alignment = Alignment(horizontal="center" if a == col else ("left" if a == col + 1 else "right"), indent=0 if a == col else 1)
        for c in range(a, b + 1):
            ws.cell(row + 1, c).fill = PatternFill("solid", fgColor="F8FAFC")
    for k, (tk, v1, f1, v2, f2) in enumerate(linhas):
        r = row + 2 + k
        for (a, b), (x, fmt, al) in zip(spans, [(k + 1, "0", "center"), (tk, None, "left"), (v1, f1, "right"), (v2, f2, "right")]):
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
            cc = ws.cell(r, a, x)
            cc.font = Font(name=F, size=9, bold=(a == col + 1), color=AZUL_S if a == col + 1 else "334155")
            cc.alignment = Alignment(horizontal=al, indent=1 if al != "center" else 0)
            if fmt:
                cc.number_format = fmt
            for c in range(a, b + 1):
                borda = ws.cell(r, c).border
                ws.cell(r, c).border = Border(left=borda.left, right=borda.right,
                                              bottom=LADO if k == len(linhas) - 1 else Side(style="thin", color=GRADE))
    return row + 2 + len(linhas)


# ------------------------------------------------------------------ dados auxiliares

def base_graficos(wb, N, D0, sub):
    """Aba oculta com as séries: proventos por mês, posição por tipo e rankings ordenados (respeitam o filtro).

    Cada ranking: coluna-chave (valor + ROW()/1E9 desempata), ticker ordenado e valores.
    Carteira vazia (no filtro): entram os tickers do catálogo com valor 0, para os gráficos nunca ficarem sem série.
    """
    ws = wb.create_sheet("Base Painel")
    ws.sheet_state = "hidden"
    AL, T = D0 + N - 1, D0 + N
    ws["A1"], ws["B1"] = "Mês", "Proventos"
    for k in range(12):
        r = 2 + k
        ws.cell(r, 1, f"=DATE(YEAR(TODAY()),MONTH(TODAY())-{11 - k},1)").number_format = "mmm/yy"
        ws.cell(r, 2, f'=SUMIFS(Proventos!$F:$F,Proventos!$A:$A,">="&A{r},Proventos!$A:$A,"<"&DATE(YEAR(A{r}),MONTH(A{r})+1,1))').number_format = K
    ws["A16"], ws["B16"] = "Tipo", "Posição"
    ws["A17"], ws["B17"] = "Ações", f"=Carteira!G{sub['AÇÃO']}"
    ws["A18"], ws["B18"] = "FIIs", f"=Carteira!G{sub['FII']}"
    ws["B17"].number_format = ws["B18"].number_format = K
    ws.cell(D0 - 1, 4, "Ticker")
    for i in range(N):
        ws.cell(D0 + i, 4, f"=Carteira!A{D0 + i}")

    specs = {   # nome: (coluna ordenada, LARGE/SMALL, [(título, coluna, formato), ...]) — tudo na aba Carteira
        "posicao": ("G", "LARGE", [("Posição", "G", K)]),
        "altas": ("I", "LARGE", [("Valorização", "I", "0.0%"), ("Lucro", "H", BRL0)]),
        "quedas": ("I", "SMALL", [("Valorização", "I", "0.0%"), ("Lucro", "H", BRL0)]),
        "meta": ("W", "LARGE", [("% da meta", "W", "0%")]),
        "proventos": ("J", "LARGE", [("Proventos", "J", K)]),
    }
    tickers = f"$D${D0}:$D${AL}"
    vazia = f'SUMIFS(Carteira!$C${D0}:$C${AL},Carteira!$AC${D0}:$AC${AL},1)=0'
    cols, c = {}, 6
    for nome, (colv, fn, valores) in specs.items():
        ck, ct = c, c + 1
        chaves = f"${L(ck)}${D0}:${L(ck)}${AL}"
        ws.cell(D0 - 1, ct, "Ticker")
        for i in range(N):
            r = D0 + i
            ws.cell(r, ck, f'=IF(AND(Carteira!AC{r}=1,OR(Carteira!C{r}>0,{vazia})),N(Carteira!{colv}{r})+ROW()/1E9,"")')
            ws.cell(r, ct, f'=IFERROR(INDEX({tickers},MATCH({fn}({chaves},{i + 1}),{chaves},0)),"")')
        out = [ct]
        for j, (tit, col_v, fmt) in enumerate(valores):
            cv = ct + 1 + j
            ws.cell(D0 - 1, cv, tit)
            for i in range(N):
                r = D0 + i
                ws.cell(r, cv, f'=IFERROR(INDEX(Carteira!${col_v}${D0}:${col_v}${AL},'
                               f'MATCH({L(ct)}{r},Carteira!$A${D0}:$A${AL},0)),"")').number_format = fmt
            out.append(cv)
        cols[nome] = tuple(out)
        c = ct + len(valores) + 2
    return ws, cols


# ------------------------------------------------------------------ painel

def painel(wb, cfg, N, D0, T, sub, classes):
    wsB, rk = base_graficos(wb, N, D0, sub)
    AL = T - 1
    ws = wb.create_sheet("Painel", 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.showRowColHeaders = False
    ws.sheet_view.zoomScale = 90
    ws.column_dimensions["A"].width = 1.5
    for c in range(ESQ0, ESQ1 + 1):
        ws.column_dimensions[L(c)].width = 6.2
    ws.column_dimensions[L(ESQ1 + 1)].width = 2
    for c in range(X0, ULT + 1):
        ws.column_dimensions[L(c)].width = 5.3
    ws.column_dimensions[L(ULT + 1)].width = 2
    FIM = 100
    _pinta(ws, 3, FIM, 1, ULT + 1, F_BG)
    cabecalho(ws)
    filtros(ws, T, D0, 47)
    ws.freeze_panes = "A4"

    C, B = "Carteira", "'Base Painel'"
    PAT = f"({C}!G{T}+SUM(OUTRAS_CLASSES))"
    NF = f"{C}!$AC${D0}:$AC${AL},1"                       # critério "no filtro"
    cols = [X0 + k * 5 for k in range(6)]
    vazia = f'SUMIFS({C}!$C${D0}:$C${AL},{NF})=0'

    # aviso de carteira vazia
    ws.merge_cells(start_row=5, start_column=X0, end_row=5, end_column=ULT)
    av = ws.cell(5, X0, f'=IF(COUNTIF({C}!C{D0}:C{AL},">0")=0,"✚  Nenhuma ação ou FII na carteira ainda — '
                        f'registre sua primeira compra em ✚ Lançar (clique aqui)","")')
    av.hyperlink = "#Cadastro!A1"
    av.font = Font(name=F, size=9, bold=True, color=LARANJA_S)
    av.alignment = Alignment(vertical="center")
    ws.row_dimensions[5].height = 24

    # cards de indicadores
    pos_t, pos_v = rk["posicao"]
    prov_t, prov_v = rk["proventos"]
    alt_t, alt_v, _ = rk["altas"]
    primeiro = lambda col: f'=IF({vazia},"—",{B}!{L(col)}{D0})'
    r = 6
    kpi(ws, r, cols[0], "PATRIMÔNIO TOTAL", f"={PAT}", BRL0,
        f'=TEXT(IF({PAT}>0,{C}!G{T}/{PAT},0),"0%")&" em renda variável"', cor_sub=SUAVE)
    kpi(ws, r, cols[1], "RENDA VARIÁVEL", f"=SUMIFS({C}!$G${D0}:$G${AL},{NF})", BRL0,
        f'=SUMIFS({C}!$C${D0}:$C${AL},{NF})', '#,##0" cotas no filtro"', SUAVE)
    kpi(ws, r, cols[2], "LUCRO / PREJUÍZO", f"=SUMIFS({C}!$H${D0}:$H${AL},{NF})", BRL0,
        f"=IF(SUMIFS({C}!$E${D0}:$E${AL},{NF})>0,SUMIFS({C}!$H${D0}:$H${AL},{NF})/SUMIFS({C}!$E${D0}:$E${AL},{NF}),0)",
        '"▲ "0.0%" sobre o investido";"▼ "0.0%" sobre o investido";"sem variação"')
    kpi(ws, r, cols[3], "PROVENTOS RECEBIDOS", f"=SUMIFS({C}!$J${D0}:$J${AL},{NF})", BRL0,
        f"=SUMIFS({C}!$V${D0}:$V${AL},{NF})", '"Renda mensal R$ "#,##0.00')
    kpi(ws, r, cols[4], "MAIOR POSIÇÃO", f'=IF({vazia},0,{B}!{L(pos_v)}{D0})', BRL0, primeiro(pos_t), cor_sub=SUAVE)
    kpi(ws, r, cols[5], "MELHOR ATIVO", primeiro(alt_t), None,
        f'=IF({vazia},"",{B}!{L(alt_v)}{D0})', '"▲ "0.0%" de valorização";"▼ "0.0%" de valorização"')
    for c in (cols[2], cols[5]):                           # verde/vermelho conforme o sinal
        sinal(ws, f"{L(c)}{r + 2}")
    sinal(ws, f"{L(cols[2])}{r + 1}", 18)

    # alocação primeiro: a rosca lê esta tabela
    r0, rt = alocacao(ws, 79, T, D0, sub, classes)

    x1, x2, x3 = X0, X0 + 10, X0 + 20
    a1, a2, a3 = L(x1), L(x2), L(x3)
    secao(ws, 10, "Carteira e proventos")
    _grafico(ws, barras(wsB, 2, wsB, 1, 2, 13, "Proventos por mês (R$)", K, AZUL_S, horizontal=False), f"{a1}11", 11, 25)
    _grafico(ws, barras(wsB, 2, wsB, 1, 17, 18, "Posição por tipo (R$)", K, LARANJA_S), f"{a2}11", 11, 25)
    _grafico(ws, barras(wsB, pos_v, wsB, pos_t, D0, AL, "Posição por ativo (R$)", K, AZUL_S), f"{a3}11", 11, 25)

    secao(ws, 27, "Alocação e meta de renda")
    _grafico(ws, rosca(ws, X0 + 5, ws, X0, r0 + 1, rt - 1, "Patrimônio por classe"), f"{a1}28", 28, 42)
    m_t, m_v = rk["meta"]
    _grafico(ws, barras(wsB, m_v, wsB, m_t, D0, AL, "Progresso da meta por ativo", "0%", LARANJA_S), f"{a2}28", 28, 42)
    _grafico(ws, barras(wsB, prov_v, wsB, prov_t, D0, AL, "Proventos por ativo (R$)", K, AZUL_S), f"{a3}28", 28, 42)

    # meta de renda
    meta = f"R$ {cfg['meta_mensal_por_ativo']:,.0f}".replace(",", ".")
    secao(ws, 44, "Meta de renda", f"cada ativo deve pagar {meta}/mês")
    n_ativos = f'COUNTIF({C}!A{D0}:A{AL},"?*")'
    r = 45
    kpi(ws, r, cols[0], "RENDA MENSAL", f"={C}!V{T}", BRL0, f"={C}!V{T}*12", '"R$ "#,##0" por ano"')
    kpi(ws, r, cols[1], "META MENSAL", f"=META_MENSAL*{n_ativos}", BRL0, "=META_MENSAL", '"R$ "#,##0" por ativo"', SUAVE)
    kpi(ws, r, cols[2], "PROGRESSO DA META", f"={C}!W{T}", PCT,
        f"=IF(META_MENSAL>0,{C}!V{T}/(META_MENSAL*{n_ativos}),0)", '0.0%" da renda total"', SUAVE)
    kpi(ws, r, cols[3], "ATIVOS NA META", f"={C}!X{T}", None, f"={C}!Y{T}", '#,##0" cotas faltantes"', SUAVE)
    kpi(ws, r, cols[4], "APORTE P/ META", f"={C}!Z{T}", BRL0, "=APORTE_MENSAL", '"Aporte de R$ "#,##0"/mês"', LARANJA_S)
    kpi(ws, r, cols[5], "TEMPO ATÉ A META", f"={C}!AA{T}", MESES,
        f"=IF(APORTE_MENSAL>0,{C}!Z{T}/APORTE_MENSAL/12,0)", '0.0" anos"', LARANJA_S)

    # destaques
    secao(ws, 49, "Destaques", "rankings no filtro escolhido")
    vpct = '"▲ "0.0%;"▼ "0.0%;0.0%'

    def top5(colunas, fmts):
        t, a, b = (L(c) for c in colunas)
        v = lambda col, k: f'=IF({vazia},"",{B}!{col}{D0 + k})'
        return [(v(t, k), v(a, k), fmts[0], v(b, k), fmts[1]) for k in range(5)]

    pos_part = [(f'=IF({vazia},"",{B}!{L(pos_t)}{D0 + k})', f'=IF({vazia},"",{B}!{L(pos_v)}{D0 + k})', BRL0,
                 f'=IF(OR({vazia},{C}!$G${T}=0),"",{B}!{L(pos_v)}{D0 + k}/{C}!$G${T})', "0.0%") for k in range(5)]
    destaque(ws, 50, x1, "Maiores posições", ["Posição", "Participação"], pos_part)
    fim = destaque(ws, 50, x2, "Maiores altas", ["Valorização", "Lucro"], top5(rk["altas"], (vpct, BRL0)))
    destaque(ws, 50, x3, "Maiores quedas", ["Valorização", "Lucro"], top5(rk["quedas"], (vpct, BRL0)))
    for c0 in (x2, x3):
        sinal(ws, f"{L(c0 + 3)}52:{L(c0 + 8)}{fim - 1}", 9)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.row_breaks.append(Break(id=43))
    return ws


def alocacao(ws, row, T, D0, sub, classes):
    """Card com a alocação por classe (valores de Carteira e Premissas). Retorna (cabeçalho, total)."""
    secao(ws, row - 1, "Alocação por classe", "valores de outras classes e pesos ideais: aba Premissas")
    r0 = row
    heads = [("Classe", 5), ("Valor (R$)", 5), ("Peso atual", 4), ("Peso ideal", 4), ("Desvio", 4), ("Aporte p/ ideal (R$)", 7)]
    spans, c = [], X0
    for h, w in heads:
        spans.append((c, c + w - 1))
        ws.merge_cells(start_row=r0, start_column=c, end_row=r0, end_column=c + w - 1)
        cc = ws.cell(r0, c, h.upper())
        cc.font = Font(name=F, size=8, bold=True, color=SUAVE)
        cc.alignment = Alignment(horizontal="left" if c == X0 else "right", vertical="center", indent=1)
        c += w
    AL = T - 1
    peso_tipo = 'PESO_RV*SUMIFS(Ativos!$F${0}:$F${1},Ativos!$B${0}:$B${1},"{2}")/Ativos!$F${3}'
    linhas = [("Ações", f"=Carteira!G{sub['AÇÃO']}", "=" + peso_tipo.format(D0, AL, "AÇÃO", T)),
              ("FIIs", f"=Carteira!G{sub['FII']}", "=" + peso_tipo.format(D0, AL, "FII", T))]
    linhas += [(f"=Premissas!A{r}", f"=Premissas!B{r}", f"=Premissas!C{r}") for r in range(classes[0], classes[1] + 1)]
    rt = r0 + 1 + len(linhas)
    card(ws, r0, rt, X0, spans[-1][1])
    for cx in range(X0, spans[-1][1] + 1):
        ws.cell(r0, cx).fill = PatternFill("solid", fgColor="F8FAFC")
    ws.row_dimensions[r0].height = 22
    V, PA, PI, DV, AP = (L(a) for a, _ in spans[1:])
    for k, (nome, val, ideal) in enumerate(linhas):
        r = r0 + 1 + k
        vals = [(nome, None, True), (val, BRL0, False), (f"=IF(${V}${rt}>0,{V}{r}/${V}${rt},0)", "0.0%", False),
                (ideal, "0.0%", False), (f"={PA}{r}-{PI}{r}", PCT, False), (f"=MAX(0,{PI}{r}*${V}${rt}-{V}{r})", BRL0, False)]
        for (a, b), (x, fmt, negrito) in zip(spans, vals):
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
            cc = ws.cell(r, a, x)
            cc.font = Font(name=F, size=9, bold=negrito, color=TXT if negrito else "334155")
            cc.alignment = Alignment(horizontal="left" if a == X0 else "right", indent=1)
            if fmt:
                cc.number_format = fmt
            for cx in range(a, b + 1):
                borda = ws.cell(r, cx).border
                ws.cell(r, cx).border = Border(left=borda.left, right=borda.right, bottom=Side(style="thin", color=GRADE))
        ws.row_dimensions[r].height = 20
    totais = [("TOTAL", None), (f"=SUM({V}{r0 + 1}:{V}{rt - 1})", BRL0), (f"=SUM({PA}{r0 + 1}:{PA}{rt - 1})", PCT),
              (f"=SUM({PI}{r0 + 1}:{PI}{rt - 1})", PCT), (None, None), (f"=SUM({AP}{r0 + 1}:{AP}{rt - 1})", BRL0)]
    for (a, b), (x, fmt) in zip(spans, totais):
        ws.merge_cells(start_row=rt, start_column=a, end_row=rt, end_column=b)
        cc = ws.cell(rt, a, x)
        cc.font = Font(name=F, size=9, bold=True, color=TOPO)
        cc.alignment = Alignment(horizontal="left" if a == X0 else "right", indent=1)
        if fmt:
            cc.number_format = fmt
        for cx in range(a, b + 1):
            ws.cell(rt, cx).fill = PatternFill("solid", fgColor="E8EEF7")
    ws.row_dimensions[rt].height = 20
    sinal(ws, f"{DV}{r0 + 1}:{DV}{rt - 1}")
    ws.print_area = f"A1:{L(ULT + 1)}{rt + 1}"
    return r0, rt
