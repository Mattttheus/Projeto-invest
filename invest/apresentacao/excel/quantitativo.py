"""Abas Quant (risco, derivadas, regressão, correlação) e Projeção (cenários com aportes).

Os números estatísticos vêm de dominio/analise.py e dominio/modelos.py (histórico semanal); a Projeção é feita com fórmulas
que partem da posição atual da Carteira e das Premissas — mudar o aporte ou as taxas recalcula tudo.
"""
import numpy as np
import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

from .abas import D0, H
from .estilo import (AZUL, BRL, BRL0, CINZA_SERIE, DATE, F, PCT, QTD, TEAL, VERDE, VERMELHO, cabecalho, cel,
                     f_bold, f_input, f_link, f_norm, f_sub, faixa, fill_head, fill_input, fill_total, fill_zebra, larguras,
                     linha_total, notas)

NUM2 = '0.00;[Red]-0.00;"-"'

# (chave em por_ativo, cabeçalho, formato)
COLUNAS = [
    ("qualidade", "Qualidade dos dados", None), ("obs", "Semanas de histórico", QTD), ("ultima", "Última cotação", DATE),
    ("ret12_preco", "Retorno 12m (preço)", PCT), ("ret12_total", "Retorno 12m (c/ proventos)", PCT),
    ("vol", "Volatilidade anual", PCT), ("sharpe", "Sharpe (vs CDI)", NUM2), ("max_queda", "Queda máx. 3a", PCT),
    ("beta", "Beta (Ibov)", NUM2), ("corr_ibov", "Correlação Ibov", NUM2),
    ("vol_ewma", "Volatilidade EWMA (λ)", PCT), ("var95", "VaR 95% semanal", PCT), ("cvar95", "CVaR 95% semanal", PCT),
    ("sortino", "Sortino", NUM2), ("capm", "Retorno esperado CAPM", PCT), ("alfa", "Alfa de Jensen", PCT),
    ("assimetria", "Assimetria", NUM2), ("curtose", "Curtose (excesso)", NUM2),
    ("jb_p", "Normalidade (p Jarque-Bera)", '0.000'), ("autocorr", "Autocorrelação 1 sem.", NUM2),
    ("velocidade", "Taxa instantânea f′(t) (anual)", PCT), ("aceleracao", "Curvatura f″(t)", NUM2),
    ("r2_poli", "R² polinômio grau 3", NUM2), ("leitura", "Leitura da tendência", None),
    ("tendencia", "Tendência anual (regressão 3a)", PCT), ("r2", "R² da regressão", NUM2),
    ("vs_tendencia", "Preço vs tendência", PCT), ("proj12", "Preço projetado 12m (R$)", BRL),
    ("proj12_min", "Faixa mín. 12m (R$)", BRL), ("proj12_max", "Faixa máx. 12m (R$)", BRL),
    ("prov12", "Proventos 12m hist. (R$)", BRL), ("cresc_prov", "Crescimento proventos a/a", PCT),
    ("peso_analise", "Peso na análise", PCT), ("contrib_risco", "Contribuição ao risco", PCT),
]


def _v(x):
    if x is None or (isinstance(x, float) and np.isnan(x)) or x is pd.NaT:
        return None
    if isinstance(x, pd.Timestamp):
        return x.to_pydatetime()
    return float(x) if isinstance(x, (np.floating, np.integer)) else x


def _sinais(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, color=VERMELHO)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=F, color=VERDE)))


def aba_quant(wb, ativos, analise, T):
    ws = wb.create_sheet("Quant")
    n = 4 + len(COLUNAS)
    if analise.get("vazio", True):
        faixa(ws, "Análise Quantitativa", "Sem histórico de cotações — rode gestao.py com internet.", 8)
        ws.cell(D0, 1, "Sem dados de histórico (dados/historico.csv).").font = f_bold
        return ws
    cart, pa = analise["carteira"], analise["por_ativo"]
    base = analise["data_base"].strftime("%d/%m/%Y")
    faixa(ws, "Análise Quantitativa", f"Histórico semanal até {base} (Yahoo Finance) • pesos: {cart['base']} "
          "• calculado por gestao.py — não digite aqui.", n)
    cabecalho(ws, H, ["Ticker", "Tipo", "Posição atual (R$)", "Peso atual"] + [c[1] for c in COLUNAS])
    ws.row_dimensions[H].height = 44
    for i, a in enumerate(ativos.itertuples()):
        r = D0 + i
        z = fill_zebra if i % 2 else None
        cel(ws, r, 1, f'=Carteira!A{r}', None, f_link, z)
        cel(ws, r, 2, f'=Carteira!B{r}', None, f_link, z)
        cel(ws, r, 3, f'=Carteira!G{r}', BRL, f_link, z)
        cel(ws, r, 4, f'=Carteira!P{r}', PCT, f_link, z)
        m = pa.loc[a.ticker] if a.ticker in pa.index else {}
        for k, (chave, _, fmt) in enumerate(COLUNAS):
            v = m.get(chave) if len(m) else None
            cel(ws, r, 5 + k, _v(v), fmt, f_norm, z)
    col = {c[0]: 5 + k for k, c in enumerate(COLUNAS)}
    linha_total(ws, T, n, "CARTEIRA")
    tot = {"ret12_total": cart.get("ret12"), "vol": cart.get("vol"), "sharpe": cart.get("sharpe"),
           "max_queda": cart.get("max_queda"), "beta": cart.get("beta"), "tendencia": cart.get("tendencia"),
           "r2": cart.get("r2"), "peso_analise": 1.0 if "vol" in cart else None,
           "contrib_risco": 1.0 if "vol" in cart else None}
    ws.cell(T, 3, f"=SUM(C{D0}:C{T - 1})").number_format = BRL
    ws.cell(T, 4, f"=SUM(D{D0}:D{T - 1})").number_format = PCT
    fmt = {c[0]: c[2] for c in COLUNAS}
    for chave, v in tot.items():
        ws.cell(T, col[chave], _v(v)).number_format = fmt[chave]

    for chave in ("ret12_preco", "ret12_total", "velocidade", "aceleracao", "tendencia", "vs_tendencia", "cresc_prov",
                  "alfa"):
        _sinais(ws, f"{L(col[chave])}{D0}:{L(col[chave])}{T}")
    q = L(col["qualidade"])
    for texto, cor in (("Alta", VERDE), ("Média", "D97706"), ("Baixa", VERMELHO)):
        ws.conditional_formatting.add(f"{q}{D0}:{q}{T - 1}", CellIsRule(operator="equal", formula=[f'"{texto}"'],
                                                                       font=Font(name=F, bold=True, color=cor)))
    c = L(col["contrib_risco"])
    ws.conditional_formatting.add(f"{c}{D0}:{c}{T - 1}", DataBarRule(start_type="num", start_value=0, end_type="max", color=TEAL))
    s = L(col["sharpe"])
    ws.conditional_formatting.add(f"{s}{D0}:{s}{T - 1}", ColorScaleRule(start_type="num", start_value=-1, start_color="FECACA",
                                                                        mid_type="num", mid_value=0, mid_color="FFFFFF",
                                                                        end_type="num", end_value=1, end_color="BBF7D0"))

    # ---- resumo da carteira (valores internos x mercado)
    r0 = T + 3
    cabecalho(ws, r0, ["Resumo da carteira", "Valor", "Leitura"])
    tir = cart.get("tir")
    cdi = analise["cfg"]["cdi"]
    resumo = [
        ("Base dos pesos", cart["base"], None, "Carteira real quando há lançamentos; senão, pesos ideais."),
        ("Retorno anual médio (3a, c/ proventos)", cart.get("ret_anual"), PCT, "Média geométrica ponderada pelos pesos."),
        ("Volatilidade anual da carteira", cart.get("vol"), PCT, "Considera as correlações entre os ativos."),
        ("Índice de diversificação", cart.get("diversificacao"), NUM2, "Vol. média dos ativos ÷ vol. da carteira (>1 = diversificação ajuda)."),
        ("Sharpe da carteira", cart.get("sharpe"), NUM2, f"Retorno acima do CDI ({cdi:.2%}) por unidade de risco."),
        ("Beta da carteira", cart.get("beta"), NUM2, "1 = oscila como o Ibovespa; <1 = mais defensiva."),
        ("Queda máxima (3a)", cart.get("max_queda"), PCT, "Pior perda do pico ao vale com os pesos atuais."),
        ("Dividend yield da carteira", cart.get("dy"), PCT, "Provento 12m ÷ cotação, ponderado."),
        ("Retorno 12m da carteira", cart.get("ret12"), PCT, "Com proventos reinvestidos."),
        ("Ibovespa 12m", cart.get("ibov12"), PCT, "Referência de mercado."),
        ("TIR realizada (seus lançamentos)", tir, PCT,
         "Taxa anual do seu dinheiro: compras, vendas, proventos e posição de hoje." if tir is not None
         else "Aparece quando houver lançamentos com data."),
        ("CDI anual (premissa)", cdi, PCT, "dados/config.json → analise.cdi_anual"),
    ]
    for k, (rot, v, fmt, dica) in enumerate(resumo):
        r = r0 + 1 + k
        cel(ws, r, 1, rot, None, f_bold)
        cel(ws, r, 2, _v(v), fmt)
        cel(ws, r, 3, dica, None, f_sub)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=10)
    _sinais(ws, f"B{r0 + 2}:B{r0 + len(resumo)}")

    # ---- matriz de correlação
    corr = analise["correlacao"]
    m0 = r0 + len(resumo) + 3
    ws.cell(m0 - 1, 1, "Correlação dos retornos semanais (3 anos) — perto de 1 = andam juntos; perto de 0 = diversificam").font = f_bold
    cabecalho(ws, m0, ["Correlação"] + list(corr.columns))
    for i, t in enumerate(corr.index):
        r = m0 + 1 + i
        cel(ws, r, 1, t, None, f_bold)
        for j, u in enumerate(corr.columns):
            cel(ws, r, 2 + j, _v(corr.loc[t, u]), NUM2)
    if len(corr):
        ult = f"{L(1 + len(corr.columns))}{m0 + len(corr)}"
        ws.conditional_formatting.add(f"B{m0 + 1}:{ult}", ColorScaleRule(
            start_type="num", start_value=-0.2, start_color="DBEAFE", mid_type="num", mid_value=0.3,
            mid_color="FFFFFF", end_type="num", end_value=1, end_color="FCA5A5"))

    notas(ws, m0 + len(corr) + 2, [
        "Derivadas: polinômio f(t) de grau 3 ajustado ao log-preço de 12 meses; f′(hoje) = taxa instantânea "
        "(e^f′ − 1 ao ano), f″(hoje) = curvatura (>0 acelerando, <0 perdendo força).",
        "EWMA: σ²ₜ = λσ²ₜ₋₁ + (1−λ)r²ₜ (λ em config.json) — pesa mais as semanas recentes. VaR/CVaR 95%: perda semanal "
        "superada em 5% das semanas / perda média nessas semanas.",
        "Jarque-Bera: p < 0,05 = retornos não normais (caudas gordas se curtose > 0). CAPM = CDI + β·(Ibov − CDI); "
        "alfa = retorno obtido − CAPM.",
        "Tendência/R²: regressão log-linear de 3 anos. R² baixo = preço sem tendência clara; não use a projeção 12m isoladamente.",
        "Faixa 12m: tendência ± 1 desvio-padrão dos resíduos (~68% dos casos históricos). Não é recomendação de compra ou venda.",
        "Contribuição ao risco: parcela da volatilidade da carteira que vem de cada ativo (soma 100%).",
        "Qualidade dos dados: Alta = 2,5+ anos de histórico, cotação recente e sem lacunas; Baixa = pouco histórico ou desatualizado.",
    ])
    larguras(ws, [30, 16, 15, 10] + [12] * len(COLUNAS))
    ws.column_dimensions[L(col["leitura"])].width = 20
    ws.freeze_panes = "B5"
    return ws


def aba_projecao(wb, analise, T):
    """Evolução mês a mês do patrimônio em ações/FIIs em 3 cenários + mesmo aporte no CDI."""
    ws = wb.create_sheet("Projeção")
    n = 9
    faixa(ws, "Projeção Futura", "Parte da posição atual (Carteira) e do aporte mensal (Premissas). "
          "Parâmetros em amarelo podem ser alterados aqui.", n)
    p = analise.get("projecao") or {"horizonte": analise["cfg"]["horizonte"], "dy": 0.06, "cdi": analise["cfg"]["cdi"],
                                    "reinvestir": analise["cfg"]["reinvestir"], "g_base": 0.0, "g_pess": -0.05, "g_otim": 0.05}
    cabecalho(ws, H, ["Parâmetro", "Valor", "Origem"])
    params = [
        ("PJ_INICIAL", "Patrimônio inicial (ações + FIIs)", f"=Carteira!G{T}", BRL0, "Carteira (posição atual)", False),
        ("PJ_APORTE", "Aporte mensal", "=APORTE_MENSAL", BRL0, "Premissas", False),
        ("PJ_DY", "Dividend yield anual", p["dy"], PCT, "Proventos 12m ÷ cotação, ponderado pela carteira", True),
        ("PJ_REINV", "Reinvestir proventos (1 = sim)", 1 if p["reinvestir"] else 0, "0", "config.json → analise", True),
        ("PJ_G_PESS", "Valorização anual — pessimista", p["g_pess"], PCT, "Base − 1 desvio-padrão no horizonte", True),
        ("PJ_G_BASE", "Valorização anual — base", p["g_base"], PCT, "Tendência 3a × R² (limitada a −10%…+20%)", True),
        ("PJ_G_OTIM", "Valorização anual — otimista", p["g_otim"], PCT, "Base + 1 desvio-padrão no horizonte", True),
        ("PJ_CDI", "CDI anual (comparação)", p["cdi"], PCT, "config.json → analise.cdi_anual", True),
        ("PJ_META", "Meta de renda mensal total", f"=SUM(Meta!G{D0}:G{T - 1})", BRL0,
         "Soma das metas de cada ativo (aba Meta)", False),
    ]
    for k, (nome, rot, v, fmt, origem, entrada) in enumerate(params):
        r = D0 + k
        cel(ws, r, 1, rot, None, f_bold)
        cel(ws, r, 2, _v(v), fmt, f_input if entrada else f_link, fill_input if entrada else None)
        cel(ws, r, 3, origem, None, f_sub)
        wb.defined_names[nome] = DefinedName(nome, attr_text=f"'Projeção'!$B${r}")

    # ---- resumo por ano (lê a tabela mensal)
    meses = int(p["horizonte"]) * 12
    t0 = D0 + len(params) + 2                          # cabeçalho da tabela mensal fica à direita (col. K)
    M0 = D0 + 1                                        # primeira linha mensal (mês 0)
    MF = M0 + meses
    r_res = t0
    cabecalho(ws, r_res, ["Ano", "Pessimista (R$)", "Base (R$)", "Otimista (R$)", "CDI (R$)",
                          "Renda mensal base (R$)", "Total aportado (R$)", "Base − CDI (R$)", "% da meta de renda"])
    for ano in range(1, int(p["horizonte"]) + 1):
        r = r_res + ano
        lin = M0 + ano * 12
        z = fill_zebra if ano % 2 else None
        cel(ws, r, 1, ano, "0", f_bold, z)
        for c, fonte in zip(range(2, 8), "MNOPQR"):
            cel(ws, r, c, f"={fonte}{lin}", BRL0, f_norm, z)
        cel(ws, r, 8, f"=C{r}-E{r}", BRL0, f_norm, z)
        cel(ws, r, 9, f"=IF(PJ_META>0,F{r}/PJ_META,0)", PCT, f_norm, z)
    fim = r_res + int(p["horizonte"])
    _sinais(ws, f"H{r_res + 1}:H{fim}")
    ws.conditional_formatting.add(f"I{r_res + 1}:I{fim}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                                     end_value=1, color=TEAL))
    r = fim + 2
    cel(ws, r, 1, "Meses até a meta de renda (base)", None, f_bold)
    cel(ws, r, 2, f'=IF(PJ_META<=0,"-",IF(COUNTIF(Q{M0}:Q{MF},"<"&PJ_META)>{meses},"além de {p["horizonte"]} anos",'
                  f'COUNTIF(Q{M0}:Q{MF},"<"&PJ_META)))', "0", f_bold, fill_total)
    cel(ws, r + 1, 1, "Renda mensal base no fim do horizonte", None, f_bold)
    cel(ws, r + 1, 2, f"=Q{MF}", BRL0, f_bold, fill_total)

    # ---- tabela mensal (colunas K..R)
    cab = ["Mês", "Data", "Pessimista (R$)", "Base (R$)", "Otimista (R$)", "CDI (R$)", "Renda mensal base (R$)",
           "Total aportado (R$)"]
    for i, h in enumerate(cab):
        c = ws.cell(M0 - 1, 11 + i, h)
        c.font, c.fill = Font(name=F, size=10, bold=True, color="FFFFFF"), fill_head
    for m in range(meses + 1):
        r = M0 + m
        ws.cell(r, 11, m).number_format = "0"
        ws.cell(r, 12, f"=EDATE(TODAY(),K{r})").number_format = "mm/yyyy"
        if m == 0:
            for c in (13, 14, 15, 16):
                ws.cell(r, c, "=PJ_INICIAL")
            ws.cell(r, 18, "=PJ_INICIAL")
        else:
            for c, g in ((13, "PJ_G_PESS"), (14, "PJ_G_BASE"), (15, "PJ_G_OTIM")):
                a = f"{L(c)}{r - 1}"
                ws.cell(r, c, f"={a}*(1+{g})^(1/12)+PJ_APORTE+PJ_REINV*{a}*PJ_DY/12")
            ws.cell(r, 16, f"=P{r - 1}*(1+PJ_CDI)^(1/12)+PJ_APORTE")
            ws.cell(r, 18, f"=R{r - 1}+PJ_APORTE")
        ws.cell(r, 17, f"=N{r}*PJ_DY/12")
        for c in range(13, 19):
            ws.cell(r, c).number_format = BRL0
        for c in range(11, 19):
            ws.cell(r, c).font = f_norm

    # ---- gráfico
    ch = LineChart()
    ch.title = "Patrimônio projetado (ações + FIIs)"
    ch.height, ch.width = 9, 22
    ch.y_axis.numFmt = ch.y_axis.number_format = '"R$ "#,##0,"k"'
    ch.y_axis.title, ch.x_axis.title = None, "Ano"
    ch.add_data(Reference(ws, min_col=2, max_col=5, min_row=r_res, max_row=fim), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=r_res + 1, max_row=fim))
    for s, cor, tr in zip(ch.series, (CINZA_SERIE, AZUL, "1E3A8A", "D97706"), (None, None, None, "dash")):
        s.graphicalProperties.line.solidFill = cor
        s.graphicalProperties.line.width = 28000
        s.smooth = False
        if tr:
            s.graphicalProperties.line.dashStyle = tr
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.legend.position = "b"
    ws.add_chart(ch, f"A{r + 4}")

    notas(ws, r + 24, [
        "Cada mês: patrimônio × (1 + valorização)^(1/12) + aporte + proventos reinvestidos (se PJ_REINV = 1).",
        "Cenários: base = tendência de 3 anos ponderada pela confiança (R²); pessimista/otimista = ± 1 desvio-padrão "
        "do retorno médio no horizonte (volatilidade ÷ √anos).",
        "CDI: mesmo aporte aplicado ao CDI da premissa, sem IR — serve de régua para o custo de oportunidade.",
        "Projeção é estimativa estatística a partir do passado; não é garantia de retorno.",
    ])
    larguras(ws, [36, 16, 16, 16, 16, 18, 17, 16, 14, 3, 7, 10, 15, 15, 15, 15, 15, 15])
    ws.freeze_panes = "A5"
    return ws


def _tabela(ws, r0, titulo, cab, linhas, fmts):
    """Título + cabeçalho + linhas [(rótulo, v1, v2, ...)]; fmts por coluna de valor. Retorna a próxima linha livre."""
    ws.cell(r0, 1, titulo).font = Font(name=F, size=11, bold=True)
    cabecalho(ws, r0 + 1, cab)
    for i, lin in enumerate(linhas):
        r = r0 + 2 + i
        z = fill_zebra if i % 2 else None
        cel(ws, r, 1, lin[0], None, f_bold, z)
        for j, v in enumerate(lin[1:]):
            cel(ws, r, 2 + j, _v(v), fmts[j] if isinstance(fmts, list) else fmts, f_norm, z)
    return r0 + 2 + len(linhas) + 2


def aba_modelos(wb, analise):
    """Álgebra linear (fatores e Markowitz), processo estocástico (Monte Carlo) e economia (Fisher)."""
    ws = wb.create_sheet("Modelos")
    mod = analise.get("modelos")
    if not mod:
        faixa(ws, "Modelos Matemáticos", "Sem histórico suficiente — rode gestao.py com internet.", 8)
        return ws
    faixa(ws, "Modelos Matemáticos", "Álgebra linear, Markowitz, Monte Carlo e economia — calculado por gestao.py "
          "a partir do histórico e da sua carteira.", 10)
    eco, fat, mc = mod["economia"], mod["fatores"], mod["monte_carlo"]
    c = analise["cfg"]

    r = _tabela(ws, 5, "1. Economia — retorno real (equação de Fisher: 1 + nominal = (1 + real)(1 + inflação))",
                ["Indicador", "Valor", "Leitura"],
                [("Inflação anual (premissa)", eco["inflacao"], "config.json → analise.inflacao_anual"),
                 ("CDI real", eco["cdi_real"], "Ganho do CDI acima da inflação"),
                 ("Retorno real da carteira (3a)", eco["carteira_real"], "Retorno médio histórico acima da inflação"),
                 ("Dividend yield real", eco["dy_real"], "Renda acima da inflação"),
                 ("Prêmio de risco sobre o CDI", eco["premio_risco"], "Quanto a renda variável pagou a mais que o CDI")],
                [PCT, None])
    _sinais(ws, "B8:B11")

    # ---- álgebra linear: decomposição espectral da covariância
    val, exp = fat["autovalores"], fat["explicado"]
    linhas = [(f"Fator {i + 1}", val[i], np.sqrt(val[i]), exp[i], exp[:i + 1].sum()) for i in range(min(5, len(val)))]
    r = _tabela(ws, r, "2. Álgebra linear — autovalores da matriz de covariância Σ = VΛVᵀ (fatores de risco); "
                f"número de condição {fat['condicao']:,.0f} (alto = ativos redundantes)",
                ["Fator", "Autovalor λ (variância)", "Volatilidade √λ", "% do risco", "% acumulado"],
                linhas, ["0.0000", PCT, PCT, PCT])
    cargas = fat["cargas"]
    r1 = r
    r = _tabela(ws, r, "Autovetores (cargas): quanto cada ativo pesa em cada fator — Fator 1 ≈ risco de mercado comum",
                ["Ativo"] + list(cargas.columns), [(t, *cargas.loc[t].values) for t in cargas.index], NUM2)
    ws.conditional_formatting.add(f"B{r1 + 2}:{L(1 + cargas.shape[1])}{r - 3}", ColorScaleRule(
        start_type="num", start_value=-0.5, start_color="DBEAFE", mid_type="num", mid_value=0, mid_color="FFFFFF",
        end_type="num", end_value=0.5, end_color="FCA5A5"))

    # ---- Markowitz
    cart, met = mod["carteiras"], mod["metricas"]
    nomes = list(cart)
    tks = [t for t in cart[nomes[0]].index if any(cart[n][t] > 0.0005 for n in nomes)]
    linhas = [(t, mod["mu"][t], *[cart[n][t] for n in nomes]) for t in tks]
    linhas += [("Retorno esperado", None, *[met[n]["ret"] for n in nomes]),
               ("Volatilidade", None, *[met[n]["vol"] for n in nomes]),
               ("Sharpe", None, *[met[n]["sharpe"] for n in nomes])]
    r1 = r
    r = _tabela(ws, r, "3. Otimização de Markowitz (sem venda a descoberto): mín. variância w ∝ Σ⁻¹1 • "
                "máx. Sharpe w ∝ Σ⁻¹(μ − CDI)", ["Ativo", "μ esperado"] + nomes, linhas, PCT)
    fim = r - 3                                               # linha do Sharpe
    for col in range(3, 3 + len(nomes)):
        ws.cell(fim, col).number_format = NUM2
        ws.conditional_formatting.add(f"{L(col)}{r1 + 2}:{L(col)}{fim - 3}", DataBarRule(
            start_type="num", start_value=0, end_type="num", end_value=0.5, color="5B8DB8"))
    for rr in range(fim - 2, fim + 1):
        for col in range(1, 3 + len(nomes)):
            ws.cell(rr, col).fill = fill_total

    # ---- Monte Carlo
    br = lambda x: f"{x:.1%}".replace(".", ",")
    ws.cell(r, 1, f"4. Monte Carlo — {mc['n']:,} trajetórias de movimento browniano geométrico ".replace(",", ".")
                  + f"(μ = {br(mc['mu'])}, σ = {br(mc['sigma'])}, aporte mensal das Premissas)").font = Font(name=F, size=11, bold=True)
    linhas = [(int(a.ano), a.p5, a.p25, a.p50, a.p75, a.p95, a.cdi, a.aportado, a.prob_cdi, a.prob_perda)
              for a in mc["anos"].itertuples()]
    r1 = r + 1
    r = _tabela(ws, r1, "Patrimônio em ações + FIIs por percentil (5% dos cenários ficam abaixo de P5)",
                ["Ano", "P5 (R$)", "P25 (R$)", "Mediana (R$)", "P75 (R$)", "P95 (R$)", "CDI (R$)", "Aportado (R$)",
                 "Prob. vencer CDI", "Prob. abaixo do aportado"], linhas, [BRL0] * 7 + [PCT, PCT])
    ano_meta = mc["ano_meta_mediano"]
    ws.cell(r - 1, 1, f"Probabilidade de atingir a meta de renda em {c['horizonte']} anos: {mc['prob_meta']:.0%}"
            + (f" • mediana: {ano_meta:.1f} anos" if ano_meta else "")).font = f_bold
    ch = LineChart()
    ch.title = "Monte Carlo: faixa de patrimônio (P5, mediana, P95) x CDI"
    ch.height, ch.width = 9, 22
    ch.y_axis.numFmt = ch.y_axis.number_format = '"R$ "#,##0,"k"'
    for col in (2, 4, 6, 7):
        ch.add_data(Reference(ws, min_col=col, min_row=r1 + 1, max_row=r - 3), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=r1 + 2, max_row=r - 3))
    for s, cor, tr in zip(ch.series, (CINZA_SERIE, AZUL, CINZA_SERIE, "D97706"), (None, None, None, "dash")):
        s.graphicalProperties.line.solidFill = cor
        s.graphicalProperties.line.width = 28000
        if tr:
            s.graphicalProperties.line.dashStyle = tr
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.legend.position = "b"
    ws.add_chart(ch, f"A{r + 1}")

    notas(ws, r + 20, [
        "μ esperado de cada ativo = 50% retorno histórico (3a) + 50% CAPM — o encolhimento reduz o erro de estimação.",
        "Fatores: se o Fator 1 explica muito do risco, a carteira depende de um único movimento de mercado.",
        "Markowitz usa o passado; trate as carteiras ótimas como referência para o peso ideal, não como ordem de compra.",
        "Monte Carlo: μ = valorização base + DY (se reinvestir); σ = volatilidade da carteira. Não é garantia de retorno.",
    ])
    larguras(ws, [30, 16, 16, 16, 16, 16, 16, 16, 15, 15])
    ws.freeze_panes = "A5"
    return ws
