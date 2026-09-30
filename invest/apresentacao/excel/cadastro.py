"""Aba Cadastro: formulário de lançamento (compra/venda de cotas) e de provento, com botões de registro.

Com macros habilitadas, o botão "Registrar" grava na hora (aba Lançamentos/Proventos + dados/*.csv) e
atualiza painel e tabelas dinâmicas — ver macros.py. Sem macros, o que ficar preenchido no formulário é
lido por ler_formulario_preenchido() e importado por aplicacao/importacao.py ao rodar atualizar.bat.
Os campos cinza vêm do cadastro do ativo (tipo, segmento, cotação atual, cotas que você tem);
preço por cota vazio = cotação atual.
"""
import warnings
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from ...dominio.carteira import OPERACOES, TIPOS_ATIVO, TIPOS_PROVENTO
from .estilo import (MODO, AZUL, BORDA_HEX, BRL, BRL0, CINZA_TXT, DATE, F, FUNDO, NAVY, QTD, TEXTO, VERDE, VERMELHO,
                           faixa, fill_input)

ABA = "Cadastro"
MARCA, MARCA_CEL = "cadastro-formulario-v3", (3, 12)
BRANCO = PatternFill("solid", fgColor="FFFFFF")
CINZA = PatternFill("solid", fgColor="F1F5F9")
BORDA = Side(style="thin", color=BORDA_HEX)
LINHA = Side(style="thin", color="EEF2F7")
FOCO = Side(style="thin", color="93C5FD")        # contorno azul-claro dos campos editáveis

# formulário: (nome definido, rótulo, formato, tipo) — "entrada" (amarelo), "auto" (cinza), "secao" (subtítulo)
TOPO = 7                                          # linha do título dos formulários
LANC_COL, PROV_COL, PAINEL_COL = 2, 5, 8          # colunas dos rótulos (campo = coluna + 1)
LANCAMENTO = [("lc_data", "Data  (vazio = hoje)", DATE, "entrada"),
              ("lc_ticker", "Ticker", None, "entrada"),
              ("lc_op", "Operação  (vazio = compra)", None, "entrada"),
              ("lc_qtd", "Quantidade de cotas", QTD, "entrada"),
              ("lc_preco", "Preço por cota  (vazio = cotação)", BRL, "entrada"),
              ("lc_custos", "Custos / taxas (R$)", BRL, "entrada"),
              (None, "DADOS DO ATIVO  ⟲ automático", None, "secao"),
              ("lc_ativo", "Ativo", None, "auto"),
              ("lc_cot", "Cotação atual (R$)", BRL, "auto"),
              ("lc_tem", "Cotas que você tem", QTD, "auto"),
              (None, "SÓ PARA TICKER NOVO", None, "secao"),
              ("lc_tipo", "Tipo do ativo (AÇÃO / FII)", None, "entrada"),
              ("lc_seg", "Segmento", None, "entrada"),
              (None, "RESUMO", None, "secao"),
              ("lc_total", "Total do lançamento", BRL, "auto"),
              ("lc_apos", "Cotas após o lançamento", QTD, "auto")]
PROVENTO = [("pv_data", "Data pagamento  (vazio = hoje)", DATE, "entrada"),
            ("pv_ticker", "Ticker", None, "entrada"),
            ("pv_tipo", "Tipo de provento", None, "entrada"),
            ("pv_valor", "Valor por cota (R$)", '"R$" #,##0.0000', "entrada"),
            ("pv_qtd", "Qtd de cotas  (vazio = suas cotas)", QTD, "entrada"),
            (None, "AUTOMÁTICO  ⟲", None, "secao"),
            ("pv_tem", "Cotas que você tem", QTD, "auto"),
            ("pv_total", "Total recebido", BRL, "auto")]
PAINEL = [("Data", DATE, "A", 11), ("Ticker", None, "B", 10), ("Operação", None, "C", 10),   # últimos lançamentos
          ("Cotas", QTD, "D", 9), ("Preço (R$)", BRL, "E", 12)]


def linha_de(campos, nome):
    """Linha do campo no formulário (os campos começam logo abaixo do título)."""
    return TOPO + 1 + [c[0] for c in campos].index(nome)


def status_linha(campos):
    return TOPO + len(campos) + 1


# ------------------------------------------------------------------ aba no Excel

def formulario(wb, ws, col, titulo, cor, campos, formulas, valores, validacoes, status):
    cr, cc = col, col + 1
    ws.merge_cells(start_row=TOPO, start_column=cr, end_row=TOPO, end_column=cc)
    for c in (cr, cc):
        ws.cell(TOPO, c).fill = PatternFill("solid", fgColor=cor)
    t = ws.cell(TOPO, cr, titulo)
    t.font = Font(name=F, size=12, bold=True, color="FFFFFF")
    t.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[TOPO].height = 30
    for i, (nome, rotulo, fmt, tipo) in enumerate(campos):
        r = TOPO + 1 + i
        if tipo == "secao":
            ws.merge_cells(start_row=r, start_column=cr, end_row=r, end_column=cc)
            s = ws.cell(r, cr, rotulo)
            s.font = Font(name=F, size=7, bold=True, color=CINZA_TXT)
            s.alignment = Alignment(vertical="bottom", indent=1)
            for c in (cr, cc):
                ws.cell(r, c).fill = BRANCO
                ws.cell(r, c).border = Border(left=BORDA if c == cr else None, right=BORDA if c == cc else None)
            ws.row_dimensions[r].height = 18
            continue
        rot = ws.cell(r, cr, rotulo)
        rot.font = Font(name=F, size=9, color=CINZA_TXT)
        rot.fill = BRANCO
        rot.alignment = Alignment(vertical="center", indent=1, shrink_to_fit=True)
        rot.border = Border(left=BORDA)
        entrada = tipo == "entrada"
        campo = ws.cell(r, cc, valores.get(nome) if entrada else formulas.get(nome))
        destaque = nome in ("lc_total", "pv_total")
        campo.fill = fill_input if entrada else CINZA
        campo.font = Font(name=F, size=12 if destaque else 10, bold=entrada or destaque,
                          color="0000FF" if entrada else (AZUL if destaque else "334155"))
        campo.alignment = Alignment(horizontal="right" if fmt else "left", vertical="center", indent=1, shrink_to_fit=True)
        lado = FOCO if entrada else BORDA
        campo.border = Border(left=lado, right=lado, top=FOCO if entrada else LINHA, bottom=FOCO if entrada else LINHA)
        if fmt:
            campo.number_format = fmt
        ws.row_dimensions[r].height = 24 if entrada else 21
        if nome:
            wb.defined_names[nome] = DefinedName(nome, attr_text=f"{ws.title}!${L(cc)}${r}")
        if nome in validacoes:
            validacoes[nome].add(campo.coordinate)
    # linha de status (mensagem do botão / da importação) + espaço dos botões
    rs = status_linha(campos)
    ws.merge_cells(start_row=rs, start_column=cr, end_row=rs, end_column=cc)
    st = ws.cell(rs, cr, valores.get("Observação"))
    st.font = Font(name=F, size=9, bold=True, color=VERMELHO if valores.get("Observação") else VERDE)
    st.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    for c in (cr, cc):
        ws.cell(rs, c).fill = BRANCO
        ws.cell(rs, c).border = Border(left=BORDA if c == cr else None, right=BORDA if c == cc else None, bottom=BORDA)
    ws.row_dimensions[rs].height = 30
    wb.defined_names[status] = DefinedName(status, attr_text=f"{ws.title}!${L(cr)}${rs}")
    ws.row_dimensions[rs + 1].height = 15            # espaço antes dos botões (altura normal: o painel ao lado usa as mesmas linhas)
    ws.row_dimensions[rs + 2].height = 32            # botões (inseridos pelo Excel em macros.py)


def lista(ws, valores, livre=False):
    dv = DataValidation(type="list", formula1=valores if valores.startswith("=") else f'"{valores}"', allow_blank=True)
    if livre:                                         # aceita ticker novo
        dv.showErrorMessage = False
    ws.add_data_validation(dv)
    return dv


def aba_cadastro(wb, faixa_tickers, n_ativos, d0, pendentes=None):
    """Formulários de lançamento e provento + 'Suas cotas hoje' + 'Últimos lançamentos'."""
    pendentes = pendentes or {"compras": [], "proventos": []}
    lc = pendentes["compras"][0] if pendentes["compras"] else {}
    pv = pendentes["proventos"][0] if pendentes["proventos"] else {}
    dl = d0 + n_ativos - 1
    ws = wb.create_sheet(ABA)
    ultima = PAINEL_COL + len(PAINEL) - 1
    faixa(ws, "Novo Lançamento", "Escolha o ticker, informe a quantidade e clique em Registrar — o resto vem do cadastro do ativo.", ultima)
    ws.sheet_view.zoomScale = 100
    ws.sheet_view.showRowColHeaders = False
    fim = status_linha(LANCAMENTO) + 4
    for r in range(4, fim + 1):
        for c in range(1, ultima + 2):
            ws.cell(r, c).fill = FUNDO
    como = ("Marque a caixa Registrar (ou menu InvestERP › Registrar)" if MODO["google"]
            else "Botão Registrar grava na hora (habilite as macros ao abrir o arquivo)")
    ajuda = ws.cell(5, LANC_COL, "Amarelo = você preenche   •   Cinza ⟲ = automático   •   " + como)
    ajuda.font = Font(name=F, size=8, italic=True, color=CINZA_TXT)
    ws.cell(*MARCA_CEL, MARCA).font = Font(color="FFFFFF", size=1)

    def busca(aba, col, tk, se_nao='""'):
        return (f'IF({tk}="","",IFERROR(INDEX({aba}!${col}${d0}:${col}${dl},'
                f'MATCH({tk},{aba}!$A${d0}:$A${dl},0)),{se_nao}))')

    tk = "lc_ticker"
    f_lc = {"lc_ativo": f'=IF({tk}="","",IFERROR(INDEX(Ativos!$B${d0}:$B${dl},MATCH({tk},Ativos!$A${d0}:$A${dl},0))'
                        f'&"  •  "&INDEX(Ativos!$C${d0}:$C${dl},MATCH({tk},Ativos!$A${d0}:$A${dl},0)),'
                        f'"Ticker novo — preencha Tipo e Segmento abaixo"))',
            "lc_cot": "=" + busca("Ativos", "D", tk),
            "lc_tem": "=" + busca("Carteira", "C", tk, "0"),
            "lc_total": '=IF(lc_qtd="","",lc_qtd*IF(lc_preco="",N(lc_cot),lc_preco)+N(lc_custos))',
            "lc_apos": '=IF(lc_qtd="","",N(lc_tem)+IF(UPPER(lc_op)="VENDA",-1,1)*lc_qtd)'}
    f_pv = {"pv_tem": "=" + busca("Carteira", "C", "pv_ticker", "0"),
            "pv_total": '=IF(pv_valor="","",pv_valor*IF(pv_qtd="",N(pv_tem),pv_qtd))'}
    v_lc = {"lc_ticker": lista(ws, f"={faixa_tickers}", livre=True), "lc_op": lista(ws, ",".join(OPERACOES)),
            "lc_tipo": lista(ws, ",".join(TIPOS_ATIVO))}
    v_pv = {"pv_ticker": lista(ws, f"={faixa_tickers}", livre=True), "pv_tipo": lista(ws, ",".join(TIPOS_PROVENTO))}
    formulario(wb, ws, LANC_COL, "✚  Compra / venda de cotas", AZUL, LANCAMENTO, f_lc, lc, v_lc, "lc_status")
    formulario(wb, ws, PROV_COL, "◈  Provento recebido", NAVY, PROVENTO, f_pv, pv, v_pv, "pv_status")

    ultimos(ws)
    if MODO["google"]:                                      # sem VBA: caixa de seleção + Apps Script (InvestERP.gs)
        for rotulo, col, nome, cor in (("✔  Registrar lançamento  →", LANC_COL, "lc_registrar", AZUL),
                                       ("✔  Registrar provento  →", PROV_COL, "pv_registrar", NAVY)):
            r = status_linha(LANCAMENTO if nome.startswith("lc") else PROVENTO) + 2
            b = ws.cell(r, col, rotulo)
            b.font = Font(name=F, size=10, bold=True, color="FFFFFF")
            b.fill = PatternFill("solid", fgColor=cor)
            b.alignment = Alignment(horizontal="center", vertical="center")
            caixa = ws.cell(r, col + 1, False)
            caixa.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            wb.defined_names[nome] = DefinedName(nome, attr_text=f"{ws.title}!${L(col + 1)}${r}")
    for col, w in ((1, 2), (2, 30), (3, 22), (4, 3), (5, 30), (6, 20), (7, 3)):
        ws.column_dimensions[L(col)].width = w
    for i, (_, _, _, w) in enumerate(PAINEL):
        ws.column_dimensions[L(PAINEL_COL + i)].width = w
    ws.print_area = f"A1:{L(ultima)}{fim}"
    ws.freeze_panes = "A4"
    return ws


def ultimos(ws, n=15):
    """Últimos lançamentos ao lado dos formulários (lidos do fim da aba Lançamentos, inclusive os do botão)."""
    c0 = PAINEL_COL
    ult = c0 + len(PAINEL) - 1
    ws.merge_cells(start_row=TOPO, start_column=c0, end_row=TOPO, end_column=ult)
    t = ws.cell(TOPO, c0, "Últimos lançamentos")
    t.font = Font(name=F, size=12, bold=True, color="FFFFFF")
    t.alignment = Alignment(vertical="center", indent=1)
    for i, (h, _, _, _) in enumerate(PAINEL):
        ws.cell(TOPO, c0 + i).fill = PatternFill("solid", fgColor="1E293B")
        c = ws.cell(TOPO + 1, c0 + i, h)
        c.font = Font(name=F, size=8, bold=True, color=CINZA_TXT)
        c.fill = PatternFill("solid", fgColor="F8FAFC")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(bottom=BORDA)
    n_lanc = "COUNTA('Lançamentos'!$B:$B)"                 # cabeçalho + lançamentos
    for k in range(n):
        r = TOPO + 2 + k
        linha = f"({n_lanc}+3-{k})"                          # k = 0 é o mais recente
        for i, (_, fmt, col, _) in enumerate(PAINEL):
            x = ws.cell(r, c0 + i, f"=IF({linha}<5,\"\",INDEX('Lançamentos'!${col}:${col},{linha}))")
            x.font = Font(name=F, size=9, bold=(col == "B"), color=TEXTO)
            x.fill = BRANCO
            x.border = Border(bottom=LINHA)
            x.alignment = Alignment(horizontal="center" if col in ("B", "C") else ("right" if fmt else "left"))
            if fmt:
                x.number_format = fmt
    nota = ws.cell(TOPO + 2 + n, c0, "Histórico completo na aba Lançamentos • posição de cada ativo na aba Carteira")
    nota.font = Font(name=F, size=8, italic=True, color=CINZA_TXT)


# ------------------------------------------------------------------ leitura do que ficou preenchido (sem macros / layouts antigos)

NOTA_ANTIGA = "Ticker novo?"                              # rodapé do formulário antigo — não é lançamento


def _limpo(v):
    if v in (None, "") or (isinstance(v, str) and (v.startswith("=") or v.startswith(NOTA_ANTIGA))):
        return None
    return v


LANC_MAPA = {"Data": "lc_data", "Ticker": "lc_ticker", "Operação": "lc_op", "Quantidade de cotas": "lc_qtd",
             "Preço por cota (R$)": "lc_preco", "Custos (R$)": "lc_custos", "Tipo do ativo": "lc_tipo",
             "Segmento": "lc_seg", "Observação": "Observação"}
PROV_MAPA = {"Data pagamento": "pv_data", "Ticker": "pv_ticker", "Tipo de provento": "pv_tipo",
             "Valor por cota (R$)": "pv_valor", "Qtd de cotas": "pv_qtd", "Observação": "Observação"}


def _ler_formulario(ws):
    """Layout atual: um formulário de lançamento e um de provento (células fixas)."""
    def ler(campos, col, mapa):
        d = {k: _limpo(ws.cell(linha_de(campos, n), col + 1).value) for k, n in mapa.items() if n != "Observação"}
        return [d] if any(v is not None for v in d.values()) else []
    return ler(LANCAMENTO, LANC_COL, LANC_MAPA), ler(PROVENTO, PROV_COL, PROV_MAPA)


def _ler_cartoes_v2(ws):
    """Layout anterior em cartões (3 por linha) — mantido para não perder nada preenchido."""
    compra = ["Data", "Ticker", "Tipo do ativo", "Segmento", "Cotação atual (R$)", "Cotas que você tem", "Operação",
              "Quantidade de cotas", "Preço por cota (R$)", "Custos (R$)"]
    provento = ["Data pagamento", "Ticker", "Tipo de provento", "Valor por cota (R$)", "Qtd de cotas"]
    digitaveis = ("Data", "Ticker", "Operação", "Quantidade de cotas", "Preço por cota (R$)", "Custos (R$)",
                  "Data pagamento", "Tipo de provento", "Valor por cota (R$)")

    def ler(campos, topo, passo, icone):
        out, k = [], 0
        while True:
            r0, c0 = topo + (k // 3) * passo, 1 + (k % 3) * 3
            if not str(ws.cell(r0, c0).value or "").startswith(icone):
                return out, k
            d = {n: _limpo(ws.cell(r0 + 1 + i, c0 + 1).value) for i, n in enumerate(campos)}
            if any(d.get(n) is not None for n in digitaveis):
                out.append(d)
            k += 1
    lc, n_c = ler(compra, 9, 15, "✚")
    pv, _ = ler(provento, 9 + -(-n_c // 3) * 15 + 2, 9, "◈")
    return lc, pv


def _ler_tabela_antiga(ws):
    """Primeiro layout (tabelas em linhas, cabeçalho na linha 8)."""
    def bloco(col0, n, renome):
        idx = {ws.cell(8, col0 + i).value: col0 + i for i in range(n)}
        out = []
        for r in range(9, ws.max_row + 1):
            d = {renome.get(h, h): _limpo(ws.cell(r, c).value) for h, c in idx.items() if h}
            if any(d.get(k) is not None for k in list(d)[:5]):
                out.append(d)
        return out
    return (bloco(1, 10, {"Quantidade": "Quantidade de cotas", "Preço (R$)": "Preço por cota (R$)"}),
            bloco(12, 7, {}))


def ler_formulario_preenchido(planilha: Path):
    """Campos preenchidos no Cadastro de uma planilha salva (layout atual ou anteriores).
    Retorna (compras, proventos, mensagem); mensagem != None quando não há o que ler."""
    if not planilha.exists():
        return [], [], "nenhuma planilha anterior"
    with warnings.catch_warnings():               # segmentação/validação estendida: irrelevantes para a leitura
        warnings.simplefilter("ignore", UserWarning)
        wb = load_workbook(planilha, data_only=False)
    if ABA not in wb.sheetnames:
        return [], [], "planilha anterior sem aba Cadastro"
    ws = wb[ABA]
    marcas = {ws.cell(3, c).value for c in (12, 13, 14)}
    if MARCA in marcas:
        return (*_ler_formulario(ws), None)
    if "cadastro-cartoes-v2" in marcas:
        return (*_ler_cartoes_v2(ws), None)
    if ws.cell(8, 1).value == "Data":
        return (*_ler_tabela_antiga(ws), None)
    return [], [], "layout do Cadastro não reconhecido — nada importado"


def para_formulario(compras, proventos):
    """Pendências (com 'Observação') no formato dos campos do formulário: uma vaga de cada tipo."""
    return {"compras": [{LANC_MAPA[k]: v for k, v in d.items() if k in LANC_MAPA} for d in compras[:1]],
            "proventos": [{PROV_MAPA[k]: v for k, v in d.items() if k in PROV_MAPA} for d in proventos[:1]]}
