"""Abas de entrada (Ativos, Lançamentos, Proventos, Premissas) e de cálculo (Carteira, Verificar).

Cada informação aparece uma única vez; as demais abas apenas apontam para ela.
Layout comum (estilo.faixa): linha 1 = faixa com marca e título, 2 = menu, 3 = subtítulo; cabeçalho na 4, dados a partir da 5.
"""
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

from ...dominio.avisos import AVISO, ERRO, INFO
from ...dominio.metas import padroes
from .estilo import (MODO, AMBAR, CINZA_TXT, BRL, BRL0, BRL4, DATE, F, MESES, PCT, QTD, TEAL, VERDE, VERMELHO,
                     cabecalho, cel, f_bold, f_input, f_link, f_norm, f_sub, faixa, fill_head2,
                     fill_input, fill_total, fill_zebra, larguras, linha_total, notas, nulo, tabela)

LIMITE = 5000         # última linha considerada em Lançamentos/Proventos
H, D0 = 4, 5          # linha do cabeçalho e primeira linha de dados (linhas 1-3: faixa, menu e subtítulo)
ORIGEM = "Gerado a partir de dados/{} — edite lá e rode gestao.py (ou atualizar.bat)."
COR_GRAVIDADE = {ERRO: VERMELHO, AVISO: AMBAR, INFO: CINZA_TXT}


def _zebra(i):
    return fill_zebra if i % 2 else None


def aba_ativos(wb, ativos, T):
    google = MODO["google"]
    ws = wb.create_sheet("Ativos")
    n = 8 if google else 7
    faixa(ws, "Cadastro de Ativos", ORIGEM.format("ativos.csv"), n)
    cabecalho(ws, H, ["Ticker", "Tipo", "Segmento", "Cotação (R$)", "Provento 12m por cota (R$)", "Peso ideal",
                      "Atualizado em"] + (["Cotação base (R$)"] if google else []))
    for i, a in enumerate(ativos.itertuples()):
        r = D0 + i
        vago = not a.ticker                                  # linha reservada a ticker novo (versão Google)
        cotacao = (f'=IF(A{r}="","",IFERROR(GOOGLEFINANCE("BVMF:"&A{r}),H{r}))' if google else a.preco)
        for col, v, fmt in ((1, a.ticker, None), (2, a.tipo, None), (3, a.segmento, None),
                            (4, cotacao, BRL), (5, a.provento_anual_cota, BRL), (6, a.peso_ideal, PCT)):
            cel(ws, r, col, None if vago and col != 4 else v, fmt, f_input, fill_input)
        quando = getattr(a, "atualizado_em", None)
        cel(ws, r, 7, None if vago else (quando if isinstance(quando, str) and quando else "manual"), None, f_sub)
        if google:                                           # último preço conhecido: reserva do GOOGLEFINANCE
            cel(ws, r, 8, None if vago else a.preco, BRL, f_input, fill_input)
    tabela(ws, "tAtivos", H, T - 1, n)
    linha_total(ws, T, n)
    ws.cell(T, 6, f"=SUM(F{D0}:F{T - 1})").number_format = PCT
    ws["E4"].comment = Comment("Soma dos proventos pagos por cota nos últimos 12 meses (Yahoo Finance). "
                               "JCP vem bruto (antes dos 15% de IR). Base da meta de renda.", "Planilha")
    notas(ws, T + 2, [
        "Cotação, provento 12m e 'Atualizado em' são preenchidos pelo Yahoo Finance a cada execução; 'manual' = não atualizado online.",
        "Peso ideal: participação desejada de cada ativo dentro da renda variável — deve somar 100%.",
    ], "Observações")
    larguras(ws, [11, 8, 26, 14, 16, 11, 14, 14])
    ws.freeze_panes = "B5"
    return ws


def aba_lancamentos(wb, lanc):
    ws = wb.create_sheet("Lançamentos")
    n = 7
    faixa(ws, "Lançamentos de Compra e Venda", ORIGEM.format("lancamentos.csv"), n)
    cabecalho(ws, H, ["Data", "Ticker", "Operação", "Quantidade", "Preço (R$)", "Custos (R$)", "Total (R$)"])
    for i, x in enumerate(lanc.itertuples()):
        r = D0 + i
        for col, v, fmt in ((1, nulo(x.data), DATE), (2, x.ticker, None), (3, x.operacao, None),
                            (4, x.quantidade, QTD), (5, x.preco, BRL), (6, x.custos, BRL)):
            cel(ws, r, col, v, fmt, f_input, _zebra(i))
        cel(ws, r, 7, f"=D{r}*E{r}+F{r}", BRL, f_norm, _zebra(i))
    ult = H + max(len(lanc), 1)
    tabela(ws, "tLancamentos", H, ult, n)
    larguras(ws, [12, 10, 11, 12, 12, 11, 15])
    ws.freeze_panes = "A5"
    return ws, ult


def aba_proventos(wb, prov):
    ws = wb.create_sheet("Proventos")
    n = 6
    faixa(ws, "Proventos Recebidos", ORIGEM.format("proventos.csv"), n)
    cabecalho(ws, H, ["Data pagamento", "Ticker", "Tipo", "Valor por cota (R$)", "Qtd de cotas", "Total (R$)"])
    for i, x in enumerate(prov.itertuples()):
        r = D0 + i
        for col, v, fmt in ((1, nulo(x.data), DATE), (2, x.ticker, None), (3, x.tipo, None),
                            (4, x.valor_cota, BRL4), (5, x.quantidade, QTD)):
            cel(ws, r, col, v, fmt, f_input, _zebra(i))
        cel(ws, r, 6, f"=D{r}*E{r}", BRL, f_norm, _zebra(i))
    ult = H + max(len(prov), 1)
    tabela(ws, "tProventos", H, ult, n)
    larguras(ws, [15, 10, 14, 16, 13, 15])
    ws.freeze_panes = "A5"
    return ws, ult


def aba_premissas(wb, cfg):
    """Parâmetros de config.json num só lugar, expostos como nomes (META_MENSAL, PRAZO_META, APORTE_MENSAL, PESO_RV)."""
    ws = wb.create_sheet("Premissas")
    n = 3
    faixa(ws, "Premissas", ORIGEM.format("config.json"), n)
    cabecalho(ws, H, ["Parâmetro", "Valor", "Uso"])
    meta_padrao, prazo_padrao = padroes(cfg)
    params = [("META_MENSAL", "Meta mensal padrão por ativo", meta_padrao, BRL0,
               "Renda que cada ativo deve pagar por mês (personalize por ativo na aba Meta)"),
              ("PRAZO_META", "Prazo padrão da meta (meses)", prazo_padrao, '0" meses"',
               "Em quanto tempo cada ativo deve atingir a meta (personalize na aba Meta)"),
              ("APORTE_MENSAL", "Aporte mensal", cfg["aporte_mensal"], BRL0, "Quanto você investe por mês (tempo até a meta)"),
              ("PESO_RV", "Peso ideal da renda variável", cfg["peso_ideal_renda_variavel"], PCT, "Ações + FIIs no patrimônio total")]
    for k, (nome, rot, v, fmt, uso) in enumerate(params):
        r = D0 + k
        cel(ws, r, 1, rot, None, f_bold)
        cel(ws, r, 2, v, fmt, f_input, fill_input)
        cel(ws, r, 3, uso, None, f_sub)
        wb.defined_names[nome] = DefinedName(nome, attr_text=f"Premissas!$B${r}")
    c0 = D0 + len(params) + 2
    cabecalho(ws, c0, ["Outras classes", "Valor (R$)", "Peso ideal no patrimônio"])
    for k, o in enumerate(cfg["outras_classes"]):
        r = c0 + 1 + k
        cel(ws, r, 1, o["classe"], None, f_bold)
        cel(ws, r, 2, o["valor"], BRL0, f_input, fill_input)
        cel(ws, r, 3, o["peso_ideal"], PCT, f_input, fill_input)
    cl = (c0 + 1, c0 + len(cfg["outras_classes"]))
    wb.defined_names["OUTRAS_CLASSES"] = DefinedName("OUTRAS_CLASSES", attr_text=f"Premissas!$B${cl[0]}:$B${cl[1]}")
    larguras(ws, [30, 16, 44])
    return ws, cl


def aba_carteira(wb, N, T, LN, PN):
    """Lista única de ações/FIIs: cadastro, posição, proventos, rebalanceamento e meta de renda na mesma linha.
    Colunas A..S mantêm a ordem usada pelo Painel e pelo Cadastro; a meta de renda fica em T..AA."""
    ws = wb.create_sheet("Carteira")
    n = 29
    faixa(ws, "Carteira", "Cada ação/FII em uma linha só: posição, proventos, peso e meta de renda — calculado, não digite aqui.", n)
    cabecalho(ws, H, ["Ticker", "Tipo", "Qtd atual", "Preço médio (R$)", "Valor investido (R$)", "Cotação (R$)",
                      "Posição atual (R$)", "Lucro / prejuízo (R$)", "Valorização", "Proventos recebidos (R$)",
                      "Proventos 12m (R$)", "Retorno total (R$)", "Retorno total %", "Yield on cost", "DY atual",
                      "Peso atual", "Peso ideal", "Desvio do peso", "Aporte p/ peso ideal (R$)",
                      "Provento 12m por cota (R$)", "Provento mensal por cota (R$)", "Renda mensal (R$)", "% da meta",
                      "Cotas p/ meta", "Cotas faltantes", "Aporte p/ meta (R$)", "Meses p/ meta", "Segmento",
                      "No filtro"])
    for c in range(20, 28):                              # bloco da meta de renda com cabeçalho diferente
        ws.cell(H, c).fill = fill_head2
    # intervalos folgados: linhas registradas pelas macros (botão Registrar) entram nos cálculos
    LN, PN = max(LN, LIMITE), max(PN, LIMITE)
    Lt, Lop, Lq, Ltot = (f"'Lançamentos'!${c}${D0}:${c}${LN}" for c in "BCDG")
    Pd, Pt, Ptot = (f"Proventos!${c}${D0}:${c}${PN}" for c in "ABF")
    fmts = {3: QTD, 4: BRL, 5: BRL, 6: BRL, 7: BRL, 8: BRL, 9: PCT, 10: BRL, 11: BRL, 12: BRL, 13: PCT,
            14: PCT, 15: PCT, 16: PCT, 17: PCT, 18: PCT, 19: BRL,
            20: BRL, 21: BRL4, 22: BRL, 23: PCT, 24: QTD, 25: QTD, 26: BRL0, 27: MESES}
    for i in range(N):
        r = D0 + i
        f = {
            1: f'=IF(Ativos!A{r}="","",Ativos!A{r})', 2: f'=IF(Ativos!B{r}="","",Ativos!B{r})',
            3: f'=SUMIFS({Lq},{Lt},A{r},{Lop},"COMPRA")-SUMIFS({Lq},{Lt},A{r},{Lop},"VENDA")',
            4: f'=IFERROR(SUMIFS({Ltot},{Lt},A{r},{Lop},"COMPRA")/SUMIFS({Lq},{Lt},A{r},{Lop},"COMPRA"),0)',
            5: f"=C{r}*D{r}", 6: f"=N(Ativos!D{r})", 7: f"=C{r}*F{r}", 8: f"=G{r}-E{r}",
            9: f"=IF(E{r}>0,H{r}/E{r},0)",
            10: f"=SUMIFS({Ptot},{Pt},A{r})",
            11: f'=SUMIFS({Ptot},{Pt},A{r},{Pd},">="&(TODAY()-365))',
            12: f"=H{r}+J{r}", 13: f"=IF(E{r}>0,L{r}/E{r},0)",
            14: f"=IF(D{r}>0,T{r}/D{r},0)", 15: f"=IF(F{r}>0,T{r}/F{r},0)",
            16: f"=IF($G${T}>0,G{r}/$G${T},0)", 17: f"=N(Ativos!F{r})", 18: f"=P{r}-Q{r}",
            19: f"=MAX(0,Q{r}*$G${T}-G{r})",
            20: f"=N(Ativos!E{r})", 21: f"=T{r}/12", 22: f"=C{r}*U{r}",
            23: f"=IF(N(Meta!G{r})>0,MIN(1,V{r}/Meta!G{r}),0)",
            24: f"=IF(U{r}>0,ROUNDUP(N(Meta!G{r})/U{r},0),0)", 25: f"=MAX(0,X{r}-C{r})",
            26: f"=Y{r}*F{r}", 27: f"=IF(APORTE_MENSAL>0,ROUNDUP(Z{r}/APORTE_MENSAL,0),0)",
            28: f'=IF(Ativos!C{r}="","",Ativos!C{r})',
            # 1 = entra nos gráficos do dashboard conforme o filtro Tipo (Todos / AÇÃO / FII) do Painel
            29: f'=IF(A{r}="",0,IF(OR(FILTRO_TIPO="Todos",B{r}=FILTRO_TIPO),1,0))',
        }
        for col, v in f.items():
            cel(ws, r, col, v, fmts.get(col), f_link if col in (1, 2, 6, 17, 20, 28) else f_norm, _zebra(i))
    tabela(ws, "tCarteira", H, T - 1, n)
    linha_total(ws, T, n)
    for col in (5, 7, 8, 10, 11, 12, 16, 17, 19, 22, 25, 26):
        ws.cell(T, col, f"=SUM({L(col)}{D0}:{L(col)}{T - 1})").number_format = PCT if col in (16, 17) else (QTD if col == 25 else BRL)
    ws.cell(T, 9, f"=IF(E{T}>0,H{T}/E{T},0)").number_format = PCT
    ws.cell(T, 13, f"=IF(E{T}>0,L{T}/E{T},0)").number_format = PCT
    ws.cell(T, 14, f"=IF(E{T}>0,SUMPRODUCT(C{D0}:C{T - 1},T{D0}:T{T - 1})/E{T},0)").number_format = PCT
    ws.cell(T, 15, f"=IF(G{T}>0,SUMPRODUCT(C{D0}:C{T - 1},T{D0}:T{T - 1})/G{T},0)").number_format = PCT
    ws.cell(T, 23, f'=IFERROR(AVERAGEIF(A{D0}:A{T - 1},"?*",W{D0}:W{T - 1}),0)').number_format = PCT
    ws.cell(T, 24, f'=COUNTIF(W{D0}:W{T - 1},">=1")&" de "&COUNTIF(A{D0}:A{T - 1},"?*")&" na meta"')
    ws.cell(T, 27, f"=IF(APORTE_MENSAL>0,ROUNDUP(Z{T}/APORTE_MENSAL,0),0)").number_format = MESES
    sub = {}
    for k, (tipo, rot) in enumerate((("AÇÃO", "Subtotal ações"), ("FII", "Subtotal FIIs"))):
        r = T + 2 + k
        sub[tipo] = r
        ws.cell(r, 1, rot).font = f_bold
        for col in (5, 7, 8, 10, 12, 22):
            cel(ws, r, col, f'=SUMIFS({L(col)}{D0}:{L(col)}{T - 1},$B${D0}:$B${T - 1},"{tipo}")', BRL, f_bold, fill_total)
        cel(ws, r, 9, f"=IF(E{r}>0,H{r}/E{r},0)", PCT, f_bold, fill_total)
        cel(ws, r, 16, f'=SUMIFS(P{D0}:P{T - 1},$B${D0}:$B${T - 1},"{tipo}")', PCT, f_bold, fill_total)
    for rng in (f"H{D0}:I{T - 1}", f"L{D0}:M{T - 1}", f"R{D0}:R{T - 1}"):
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, color=VERMELHO)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=F, color=VERDE)))
    ws.conditional_formatting.add(f"P{D0}:P{T - 1}", DataBarRule(start_type="num", start_value=0, end_type="max", color="5B8DB8"))
    ws.conditional_formatting.add(f"W{D0}:W{T - 1}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=TEAL))
    # ativos do catálogo que você ainda não tem ficam em cinza (a lista é uma só: catálogo + carteira)
    ws.conditional_formatting.add(f"A{D0}:AB{T - 1}", FormulaRule(formula=[f"$C{D0}=0"], font=Font(name=F, color="CBD5E1")))
    notas(ws, T + 5, [
        "Linhas em cinza: ativos do catálogo que você ainda não tem (a mesma lista serve de catálogo e de carteira).",
        "Meta de renda: cotas p/ meta = meta mensal ÷ provento mensal por cota; aporte = cotas faltantes × cotação.",
        "Tipo, segmento, cotação, provento 12m e peso ideal vêm de dados/ativos.csv (aba Ativos, oculta).",
    ])
    larguras(ws, [10, 7, 10, 12, 15, 11, 15, 15, 11, 14, 13, 14, 11, 10, 9, 9, 9, 10, 15,
                  13, 13, 13, 10, 11, 11, 14, 12, 22, 8])
    ws.column_dimensions["AC"].hidden = True             # auxiliar do filtro do dashboard
    ws.freeze_panes = "C5"
    return ws, sub


def aba_verificar(wb, avisos):
    ws = wb.create_sheet("Verificar")
    n = 3
    faixa(ws, "Pontos a Conferir", "Checagens automáticas feitas a cada geração — corrija em dados/.", n)
    cabecalho(ws, H, ["Gravidade", "Ativo", "O que foi encontrado"])
    if not avisos:
        ws.cell(D0, 1, "Nenhuma inconsistência encontrada.").font = f_bold
    for k, (g, t, m) in enumerate(avisos):
        r = D0 + k
        cel(ws, r, 1, g, None, Font(name=F, size=10, bold=True, color=COR_GRAVIDADE.get(g, CINZA_TXT)))
        cel(ws, r, 2, t, None, f_bold)
        for c in (1, 2):
            ws.cell(r, c).alignment = Alignment(vertical="center", indent=1)
        cel(ws, r, 3, m).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 30
    tabela(ws, "tVerificar", H, H + len(avisos), n)
    larguras(ws, [11, 24, 100])
    return ws
