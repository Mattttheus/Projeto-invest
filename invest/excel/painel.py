"""Aba Painel: dashboard escuro (estilo "Dashboard de Vendas"), no Excel e no Google Planilhas.

Layout: barra de navegação (linha 1) | cabeçalho "DASHBOARD INVESTIMENTOS" com filtro Tipo e botões
| trilho de ícones à esquerda | faixa de indicadores com ícones | 6 gráficos em painéis escuros
(colunas verdes, barras amarelas, colunas laranja, rosca) | meta de renda | destaques | alocação.
Só lê as outras abas; o filtro Tipo (Todos / AÇÃO / FII) age sobre os rankings via Carteira!AC.
"""
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

from .estilo import BRL0, F, MESES, PCT, SIDEBAR, MODO, navbar_texto

# ---- paleta do dashboard escuro
BG, PAINEL, BORDA_ESC, LINHA_ESC = "0B1F3A", "12294A", "1F3B63", "1B355C"
BRANCO_T, SUAVE, APAGADO = "FFFFFF", "CBD5E1", "7C8DA6"
VERDE_G, AMARELO, LARANJA, AZUL_G, VERMELHO_G = "1FB585", "F5C518", "F28C28", "3B82F6", "EF4444"
POSITIVO, NEGATIVO = "34D399", "F87171"
ROSCA = [VERDE_G, AMARELO, LARANJA, AZUL_G, VERMELHO_G, "A78BFA"]

# grade: A = trilho de ícones | B = respiro | C..AE = conteúdo (29 colunas) | AF = margem
X0 = 3
ULT = X0 + 28
W3 = 9.9                                             # largura (cm) de cada gráfico (3 por linha)
K = '"R$ "#,##0,"k";-"R$ "#,##0,"k";"R$ 0"'          # valores compactos: R$ 40k
F_BG, F_PAINEL = PatternFill("solid", fgColor=BG), PatternFill("solid", fgColor=PAINEL)
BORDA = Side(style="thin", color=BORDA_ESC)
TRILHO = [("⌂", "Painel"), ("◧", "Carteira"), ("▦", "Análises"), ("✚", "Cadastro"), ("⇄", "Lançamentos"),
          ("◈", "Proventos"), ("⚙", "Premissas"), ("⚠", "Verificar")]


# ------------------------------------------------------------------ gráficos escuros

def _texto(tamanho=800, cor=SUAVE, negrito=False):
    cp = CharacterProperties(sz=tamanho, b=negrito, solidFill=cor)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _titulo(texto):
    cp = CharacterProperties(sz=1100, b=True, solidFill=BRANCO_T)
    rico = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=texto)])])
    return Title(tx=Text(rich=rico), overlay=False)


def _fundo(ch):
    ch.graphical_properties = GraphicalProperties(solidFill=PAINEL, ln=LineProperties(solidFill=BORDA_ESC))
    ch.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))


def barras(ws_dados, col, ws_cat, col_cat, r1, r2, titulo, fmt, cor, horizontal=True):
    """Barras/colunas de uma cor, com valor na ponta, em painel escuro."""
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.add_data(Reference(ws_dados, min_col=col, min_row=r1 - 1, max_row=r2), titles_from_data=True)
    ch.set_categories(Reference(ws_cat, min_col=col_cat, min_row=r1, max_row=r2))
    ch.y_axis.numFmt = ch.y_axis.number_format = fmt
    ch.gapWidth = 45
    ch.x_axis.tickLblSkip = 1
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"        # maior no topo
        ch.y_axis.crosses = "max"
        ch.x_axis.tickLblPos = "low"
    ch.legend = None
    ch.dataLabels = DataLabelList(showVal=True, showLegendKey=False, showCatName=False, showSerName=False,
                                  showPercent=False, dLblPos="outEnd")
    ch.dataLabels.txPr = _texto(800, BRANCO_T)
    ch.title = _titulo(titulo)
    ch.style = 2
    ch.height, ch.width = 7.4, W3
    _fundo(ch)
    s = ch.series[0]
    s.graphicalProperties.solidFill = cor
    s.graphicalProperties.line.solidFill = cor
    s.invertIfNegative = False
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.x_axis.txPr = ch.y_axis.txPr = _texto(800, SUAVE)
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINHA_ESC))
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINHA_ESC))
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    return ch


def rosca(ws_dados, col, ws_cat, col_cat, r1, r2, titulo):
    """Rosca colorida com % em cada fatia (como 'Quantidade faturada por região')."""
    ch = DoughnutChart()
    ch.holeSize = 58
    ch.add_data(Reference(ws_dados, min_col=col, min_row=r1 - 1, max_row=r2), titles_from_data=True)
    ch.set_categories(Reference(ws_cat, min_col=col_cat, min_row=r1, max_row=r2))
    ch.dataLabels = DataLabelList(showPercent=True, showVal=False, showCatName=False, showSerName=False,
                                  showLegendKey=False, showLeaderLines=False)
    ch.dataLabels.txPr = _texto(800, BRANCO_T, True)
    ch.legend.position = "l"
    ch.legend.txPr = _texto(800, SUAVE)
    for i in range(r2 - r1 + 1):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = ROSCA[i % len(ROSCA)]
        pt.graphicalProperties.line.solidFill = PAINEL
        ch.series[0].dPt.append(pt)
    ch.title = _titulo(titulo)
    ch.style = 2
    ch.height, ch.width = 7.4, W3
    _fundo(ch)
    return ch


# ------------------------------------------------------------------ blocos da tela

def _pinta(ws, r1, r2, c1, c2, fill):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).fill = fill


def cabecalho(ws, D0, AL):
    """Linha 1 = navegação; linhas 2-3 = título grande, filtro Tipo (lista suspensa) e espaço dos botões."""
    for c in range(1, ULT + 2):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=SIDEBAR)
    ws.row_dimensions[1].height = 30
    if MODO["google"]:
        navbar_texto(ws, ULT + 1)
    else:
        m = ws.cell(1, 1, "  ◆ InvestERP")
        m.font = Font(name=F, size=12, bold=True, color=BRANCO_T)
        m.alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 26
    ws.merge_cells(start_row=2, start_column=X0, end_row=3, end_column=X0 + 12)
    t = ws.cell(2, X0, "DASHBOARD INVESTIMENTOS")
    t.font = Font(name=F, size=22, bold=True, color=BRANCO_T)
    t.alignment = Alignment(vertical="center")
    # filtro Tipo: rótulo + caixa branca com lista (como os filtros Data/Estado/Produto da imagem)
    ws.merge_cells(start_row=2, start_column=X0 + 14, end_row=2, end_column=X0 + 18)
    rot = ws.cell(2, X0 + 14, "Tipo de ativo")
    rot.font = Font(name=F, size=9, bold=True, color=BRANCO_T)
    rot.alignment = Alignment(vertical="bottom")
    ws.merge_cells(start_row=3, start_column=X0 + 14, end_row=3, end_column=X0 + 18)
    filtro = ws.cell(3, X0 + 14, "Todos")
    filtro.font = Font(name=F, size=10, bold=True, color="0F172A")
    filtro.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in range(X0 + 14, X0 + 19):
        ws.cell(3, c).fill = PatternFill("solid", fgColor="FFFFFF")
    dv = DataValidation(type="list", formula1='"Todos,AÇÃO,FII"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(filtro.coordinate)
    ws.parent.defined_names["FILTRO_TIPO"] = DefinedName("FILTRO_TIPO", attr_text=f"Painel!${L(X0 + 14)}$3")
    fonte = "Google Finance" if MODO["google"] else "Yahoo Finance"
    sub = ws.cell(4, X0, "=TODAY()")                        # data pelo formato da célula (vale em qualquer idioma)
    sub.number_format = f'"Atualizado em "dd/mm/yyyy"   •   cotações: {fonte}   •   use o filtro Tipo ao lado"'
    sub.font = Font(name=F, size=8, color=APAGADO)
    sub.alignment = Alignment(horizontal="left")
    ws.row_dimensions[4].height = 16


def trilho(ws, ate):
    """Coluna A: ícones com link para as abas (como a barra lateral de ícones da imagem)."""
    for r in range(2, ate + 1):
        ws.cell(r, 1).fill = PatternFill("solid", fgColor="081629")
    for k, (icone, aba) in enumerate(TRILHO):
        r = 6 + k * 2
        c = ws.cell(r, 1, icone)
        c.font = Font(name="Segoe UI Symbol", size=15, bold=True, color=AMARELO if aba == "Painel" else SUAVE)
        c.alignment = Alignment(horizontal="center", vertical="center")
        if aba != "Painel":
            c.hyperlink = f"#'{aba}'!A1" if not aba.isalnum() else f"#{aba}!A1"


def kpi(ws, row, col, icone, rotulo, valor, fmt, sub, sub_fmt=None, cor_icone=AMARELO, span=4):
    """Indicador da faixa escura: ícone grande + valor em destaque + rótulo, como 'Faturamento / Maior Venda'."""
    _pinta(ws, row, row + 2, col, col + span - 1, F_PAINEL)
    for c in range(col, col + span):
        ws.cell(row, c).border = Border(top=BORDA)
        ws.cell(row + 2, c).border = Border(bottom=BORDA)
    ws.merge_cells(start_row=row, start_column=col, end_row=row + 2, end_column=col)
    ic = ws.cell(row, col, icone)
    ic.font = Font(name="Segoe UI Emoji", size=20, color=cor_icone)
    ic.alignment = Alignment(horizontal="center", vertical="center")
    for dr, (x, fnt, fm) in enumerate(((rotulo, Font(name=F, size=8, bold=True, color=SUAVE), None),
                                       (valor, Font(name=F, size=17, bold=True, color=BRANCO_T), fmt),
                                       (sub, Font(name=F, size=8, color=AMARELO), sub_fmt))):
        ws.merge_cells(start_row=row + dr, start_column=col + 1, end_row=row + dr, end_column=col + span - 1)
        c = ws.cell(row + dr, col + 1, x)
        c.font = fnt
        c.alignment = Alignment(vertical="center", indent=1, shrink_to_fit=True)
        if fm:
            c.number_format = fm
    ws.row_dimensions[row + 1].height = 30
    return ws.cell(row + 1, col + 1)


def secao(ws, row, texto, dica=""):
    c = ws.cell(row, X0, texto)
    c.font = Font(name=F, size=11, bold=True, color=BRANCO_T)
    if dica:
        d = ws.cell(row, X0 + 12, dica)
        d.font = Font(name=F, size=8, color=APAGADO)
    ws.row_dimensions[row].height = 20


def painel_grafico(ws, row, col, altura=15):
    """Moldura escura sob cada gráfico (o gráfico já é escuro; isto preenche as frestas)."""
    _pinta(ws, row, row + altura - 1, col, col + 8, F_PAINEL)


def sinal(ws, faixa, tamanho=None):
    for op, cor in (("lessThan", NEGATIVO), ("greaterThan", POSITIVO)):
        ws.conditional_formatting.add(faixa, CellIsRule(operator=op, formula=["0"],
                                                        font=Font(name=F, size=tamanho, bold=bool(tamanho), color=cor)))


def destaque(ws, row, col, titulo, cabecalhos, linhas):
    """Mini-tabela escura de 9 colunas: #, ticker e dois valores."""
    spans = [(col, col), (col + 1, col + 2), (col + 3, col + 5), (col + 6, col + 8)]
    _pinta(ws, row, row + 1 + len(linhas), col, col + 8, F_PAINEL)
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 8)
    t = ws.cell(row, col, titulo)
    t.font = Font(name=F, size=10, bold=True, color=BRANCO_T)
    t.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[row].height = 22
    for (a, b), h in zip(spans, ["#", "Ativo", *cabecalhos]):
        ws.merge_cells(start_row=row + 1, start_column=a, end_row=row + 1, end_column=b)
        cc = ws.cell(row + 1, a, h.upper())
        cc.font = Font(name=F, size=7, bold=True, color=APAGADO)
        cc.alignment = Alignment(horizontal="center" if a == col else ("left" if a == col + 1 else "right"), indent=0 if a == col else 1)
        for c in range(a, b + 1):
            ws.cell(row + 1, c).border = Border(bottom=BORDA)
    for k, (tk, v1, f1, v2, f2) in enumerate(linhas):
        r = row + 2 + k
        for (a, b), (x, fmt, al) in zip(spans, [(k + 1, "0", "center"), (tk, None, "left"), (v1, f1, "right"), (v2, f2, "right")]):
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
            cc = ws.cell(r, a, x)
            cc.font = Font(name=F, size=9, bold=(a == col + 1), color=AMARELO if a == col + 1 else SUAVE)
            cc.alignment = Alignment(horizontal=al, indent=1 if al != "center" else 0)
            if fmt:
                cc.number_format = fmt
            for c in range(a, b + 1):
                ws.cell(r, c).border = Border(bottom=Side(style="thin", color=LINHA_ESC))
    return row + 2 + len(linhas)


# ------------------------------------------------------------------ dados auxiliares

def base_graficos(wb, N, D0, sub):
    """Aba oculta com as séries: proventos por mês, posição por tipo e rankings ordenados (respeitam o filtro).

    Cada ranking: coluna-chave (valor + ROW()/1E9 desempata), ticker ordenado e valores.
    Carteira vazia (no filtro): entram os tickers do catálogo com valor 0, para os gráficos nunca ficarem sem série.
    """
    ws = wb.create_sheet("Base Painel")
    ws.sheet_state = "hidden"
    AL, T = D0 + N - 1, D0 + N
    ws["A1"], ws["B1"] = "Mês", "Proventos"
    for k in range(12):
        r = 2 + k
        ws.cell(r, 1, f"=DATE(YEAR(TODAY()),MONTH(TODAY())-{11 - k},1)").number_format = "mmm/yy"
        ws.cell(r, 2, f'=SUMIFS(Proventos!$F:$F,Proventos!$A:$A,">="&A{r},Proventos!$A:$A,"<"&DATE(YEAR(A{r}),MONTH(A{r})+1,1))').number_format = K
    ws["A16"], ws["B16"] = "Tipo", "Posição"
    ws["A17"], ws["B17"] = "Ações", f"=Carteira!G{sub['AÇÃO']}"
    ws["A18"], ws["B18"] = "FIIs", f"=Carteira!G{sub['FII']}"
    ws["B17"].number_format = ws["B18"].number_format = K
    ws.cell(D0 - 1, 4, "Ticker")
    for i in range(N):
        ws.cell(D0 + i, 4, f"=Carteira!A{D0 + i}")

    specs = {   # nome: (coluna ordenada, LARGE/SMALL, [(título, coluna, formato), ...]) — tudo na aba Carteira
        "posicao": ("G", "LARGE", [("Posição", "G", K)]),
        "altas": ("I", "LARGE", [("Valorização", "I", "0.0%"), ("Lucro", "H", BRL0)]),
        "quedas": ("I", "SMALL", [("Valorização", "I", "0.0%"), ("Lucro", "H", BRL0)]),
        "meta": ("W", "LARGE", [("% da meta", "W", "0%")]),
        "proventos": ("J", "LARGE", [("Proventos", "J", K)]),
    }
    tickers = f"$D${D0}:$D${AL}"
    vazia = f'SUMIFS(Carteira!$C${D0}:$C${AL},Carteira!$AC${D0}:$AC${AL},1)=0'
    cols, c = {}, 6
    for nome, (colv, fn, valores) in specs.items():
        ck, ct = c, c + 1
        chaves = f"${L(ck)}${D0}:${L(ck)}${AL}"
        ws.cell(D0 - 1, ct, "Ticker")
        for i in range(N):
            r = D0 + i
            ws.cell(r, ck, f'=IF(AND(Carteira!AC{r}=1,OR(Carteira!C{r}>0,{vazia})),N(Carteira!{colv}{r})+ROW()/1E9,"")')
            ws.cell(r, ct, f'=IFERROR(INDEX({tickers},MATCH({fn}({chaves},{i + 1}),{chaves},0)),"")')
        out = [ct]
        for j, (tit, col_v, fmt) in enumerate(valores):
            cv = ct + 1 + j
            ws.cell(D0 - 1, cv, tit)
            for i in range(N):
                r = D0 + i
                ws.cell(r, cv, f'=IFERROR(INDEX(Carteira!${col_v}${D0}:${col_v}${AL},'
                               f'MATCH({L(ct)}{r},Carteira!$A${D0}:$A${AL},0)),"")').number_format = fmt
            out.append(cv)
        cols[nome] = tuple(out)
        c = ct + len(valores) + 2
    return ws, cols


# ------------------------------------------------------------------ painel

def painel(wb, wsC, cfg, N, D0, T, sub, classes):
    wsB, rk = base_graficos(wb, N, D0, sub)
    AL = T - 1
    ws = wb.create_sheet("Painel", 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.showRowColHeaders = False
    ws.sheet_view.zoomScale = 90
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 1.5
    for c in range(X0, ULT + 1):
        ws.column_dimensions[L(c)].width = 5.3
    ws.column_dimensions[L(ULT + 1)].width = 2
    FIM = 100
    _pinta(ws, 2, FIM, 1, ULT + 1, F_BG)
    cabecalho(ws, D0, AL)
    trilho(ws, FIM)
    ws.freeze_panes = "B5"

    C, B = "Carteira", "'Base Painel'"
    PAT = f"({C}!G{T}+SUM(OUTRAS_CLASSES))"
    NF = f"{C}!$AC${D0}:$AC${AL},1"                       # critério "no filtro"
    cols = [X0 + k * 5 for k in range(6)]
    vazia = f'SUMIFS({C}!$C${D0}:$C${AL},{NF})=0'

    # aviso de carteira vazia
    ws.merge_cells(start_row=5, start_column=X0, end_row=5, end_column=ULT)
    av = ws.cell(5, X0, f'=IF(COUNTIF({C}!C{D0}:C{AL},">0")=0,"✚  Nenhuma ação ou FII na carteira ainda — '
                        f'registre sua primeira compra em ✚ Lançar (clique aqui)","")')
    av.hyperlink = "#Cadastro!A1"
    av.font = Font(name=F, size=9, bold=True, color=AMARELO)
    ws.row_dimensions[5].height = 18

    # faixa de indicadores (como Faturamento / Qtd faturada / Maior venda / Maior receita / Melhor vendedor)
    pos_t, pos_v = rk["posicao"]
    prov_t, prov_v = rk["proventos"]
    alt_t, alt_v, _ = rk["altas"]
    primeiro = lambda col: f'=IF({vazia},"—",{B}!{L(col)}{D0})'
    r = 6
    kpi(ws, r, cols[0], "💰", "PATRIMÔNIO TOTAL", f"={PAT}", BRL0,
        f'=TEXT(IF({PAT}>0,{C}!G{T}/{PAT},0),"0%")&" em renda variável"')
    kpi(ws, r, cols[1], "📈", "RENDA VARIÁVEL", f"=SUMIFS({C}!$G${D0}:$G${AL},{NF})", BRL0,
        f'=SUMIFS({C}!$C${D0}:$C${AL},{NF})', '#,##0" cotas no filtro"')
    kpi(ws, r, cols[2], "🪙", "PROVENTOS RECEBIDOS", f"=SUMIFS({C}!$J${D0}:$J${AL},{NF})", BRL0,
        f"=SUMIFS({C}!$V${D0}:$V${AL},{NF})", '"Renda mensal R$ "#,##0.00')
    kpi(ws, r, cols[3], "🏦", "MAIOR POSIÇÃO", f'=IF({vazia},0,{B}!{L(pos_v)}{D0})', BRL0, primeiro(pos_t))
    kpi(ws, r, cols[4], "💵", "MAIOR PROVENTO", f'=IF({vazia},0,{B}!{L(prov_v)}{D0})', BRL0, primeiro(prov_t))
    kpi(ws, r, cols[5], "🏅", "MELHOR ATIVO", primeiro(alt_t), None,
        f'=IF({vazia},"",{B}!{L(alt_v)}{D0})', '"▲ "0.0%" de valorização";"▼ "0.0%" de valorização"')

    # alocação primeiro: a rosca lê esta tabela
    r0, rt = alocacao(ws, 79, T, D0, sub, classes)

    x1, x2, x3 = X0, X0 + 10, X0 + 20
    a1, a2, a3 = L(x1), L(x2), L(x3)
    secao(ws, 10, "Carteira e proventos")
    for x in (x1, x2, x3):
        painel_grafico(ws, 11, x)
    ws.add_chart(barras(wsB, 2, wsB, 1, 2, 13, "Proventos por mês (R$)", K, VERDE_G, horizontal=False), f"{a1}11")
    ws.add_chart(barras(wsB, 2, wsB, 1, 17, 18, "Posição por tipo (R$)", K, AMARELO), f"{a2}11")
    ws.add_chart(barras(wsB, pos_v, wsB, pos_t, D0, AL, "Posição por ativo (R$)", K, LARANJA, horizontal=False), f"{a3}11")

    secao(ws, 27, "Alocação e meta de renda")
    for x in (x1, x2, x3):
        painel_grafico(ws, 28, x)
    ws.add_chart(rosca(ws, X0 + 5, ws, X0, r0 + 1, rt - 1, "Patrimônio por classe"), f"{a1}28")
    m_t, m_v = rk["meta"]
    ws.add_chart(barras(wsB, m_v, wsB, m_t, D0, AL, "Progresso da meta por ativo", "0%", AMARELO), f"{a2}28")
    ws.add_chart(barras(wsB, prov_v, wsB, prov_t, D0, AL, "Proventos por ativo (R$)", K, LARANJA, horizontal=False), f"{a3}28")

    # meta de renda (faixa de indicadores)
    meta = f"R$ {cfg['meta_mensal_por_ativo']:,.0f}".replace(",", ".")
    secao(ws, 44, "Meta de renda", f"cada ativo deve pagar {meta}/mês")
    n_ativos = f'COUNTIF({C}!A{D0}:A{AL},"?*")'
    r = 45
    kpi(ws, r, cols[0], "🎯", "RENDA MENSAL", f"={C}!V{T}", BRL0, f"={C}!V{T}*12", '"R$ "#,##0" por ano"', VERDE_G)
    kpi(ws, r, cols[1], "🧭", "META MENSAL", f"=META_MENSAL*{n_ativos}", BRL0, "=META_MENSAL", '"R$ "#,##0" por ativo"', VERDE_G)
    kpi(ws, r, cols[2], "📊", "PROGRESSO DA META", f"={C}!W{T}", PCT,
        f"=IF(META_MENSAL>0,{C}!V{T}/(META_MENSAL*{n_ativos}),0)", '0.0%" da renda total"', VERDE_G)
    kpi(ws, r, cols[3], "✅", "ATIVOS NA META", f"={C}!X{T}", None, f"={C}!Y{T}", '#,##0" cotas faltantes"', VERDE_G)
    kpi(ws, r, cols[4], "💸", "APORTE P/ META", f"={C}!Z{T}", BRL0, "=APORTE_MENSAL", '"Aporte de R$ "#,##0"/mês"', LARANJA)
    kpi(ws, r, cols[5], "⏳", "TEMPO ATÉ A META", f"={C}!AA{T}", MESES,
        f"=IF(APORTE_MENSAL>0,{C}!Z{T}/APORTE_MENSAL/12,0)", '0.0" anos"', LARANJA)

    # destaques
    secao(ws, 49, "Destaques", "rankings no filtro escolhido")
    vpct = '"▲ "0.0%;"▼ "0.0%;0.0%'

    def top5(colunas, fmts):
        t, a, b = (L(c) for c in colunas)
        v = lambda col, k: f'=IF({vazia},"",{B}!{col}{D0 + k})'
        return [(v(t, k), v(a, k), fmts[0], v(b, k), fmts[1]) for k in range(5)]

    pos_part = [(f'=IF({vazia},"",{B}!{L(pos_t)}{D0 + k})', f'=IF({vazia},"",{B}!{L(pos_v)}{D0 + k})', BRL0,
                 f'=IF(OR({vazia},{C}!$G${T}=0),"",{B}!{L(pos_v)}{D0 + k}/{C}!$G${T})', "0.0%") for k in range(5)]
    destaque(ws, 50, x1, "🏦  Maiores posições", ["Posição", "Participação"], pos_part)
    fim = destaque(ws, 50, x2, "📈  Maiores altas", ["Valorização", "Lucro"], top5(rk["altas"], (vpct, BRL0)))
    destaque(ws, 50, x3, "📉  Maiores quedas", ["Valorização", "Lucro"], top5(rk["quedas"], (vpct, BRL0)))
    for c0 in (x2, x3):
        sinal(ws, f"{L(c0 + 3)}52:{L(c0 + 8)}{fim - 1}", 9)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.row_breaks.append(Break(id=43))
    return ws


def alocacao(ws, row, T, D0, sub, classes):
    """Tabela escura de alocação por classe (valores de Carteira e Premissas). Retorna (cabeçalho, total)."""
    secao(ws, row - 1, "Alocação por classe", "valores de outras classes e pesos ideais: aba Premissas")
    r0 = row
    heads = [("Classe", 5), ("Valor (R$)", 5), ("Peso atual", 4), ("Peso ideal", 4), ("Desvio", 4), ("Aporte p/ ideal (R$)", 7)]
    spans, c = [], X0
    for h, w in heads:
        spans.append((c, c + w - 1))
        ws.merge_cells(start_row=r0, start_column=c, end_row=r0, end_column=c + w - 1)
        cc = ws.cell(r0, c, h)
        cc.font = Font(name=F, size=9, bold=True, color=BRANCO_T)
        cc.alignment = Alignment(horizontal="center", vertical="center")
        for cx in range(c, c + w):
            ws.cell(r0, cx).fill = PatternFill("solid", fgColor=BORDA_ESC)
        c += w
    ws.row_dimensions[r0].height = 22
    AL = T - 1
    peso_tipo = 'PESO_RV*SUMIFS(Ativos!$F${0}:$F${1},Ativos!$B${0}:$B${1},"{2}")/Ativos!$F${3}'
    linhas = [("Ações", f"=Carteira!G{sub['AÇÃO']}", "=" + peso_tipo.format(D0, AL, "AÇÃO", T)),
              ("FIIs", f"=Carteira!G{sub['FII']}", "=" + peso_tipo.format(D0, AL, "FII", T))]
    linhas += [(f"=Premissas!A{r}", f"=Premissas!B{r}", f"=Premissas!C{r}") for r in range(classes[0], classes[1] + 1)]
    rt = r0 + 1 + len(linhas)
    V, PA, PI, DV, AP = (L(a) for a, _ in spans[1:])
    for k, (nome, val, ideal) in enumerate(linhas):
        r = r0 + 1 + k
        vals = [(nome, None, True), (val, BRL0, False), (f"=IF(${V}${rt}>0,{V}{r}/${V}${rt},0)", "0.0%", False),
                (ideal, "0.0%", False), (f"={PA}{r}-{PI}{r}", PCT, False), (f"=MAX(0,{PI}{r}*${V}${rt}-{V}{r})", BRL0, False)]
        for (a, b), (x, fmt, negrito) in zip(spans, vals):
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
            cc = ws.cell(r, a, x)
            cc.font = Font(name=F, size=9, bold=negrito, color=BRANCO_T if negrito else SUAVE)
            cc.alignment = Alignment(horizontal="left" if a == X0 else "right", indent=1)
            if fmt:
                cc.number_format = fmt
            for cx in range(a, b + 1):
                ws.cell(r, cx).fill = F_PAINEL
                ws.cell(r, cx).border = Border(bottom=Side(style="thin", color=LINHA_ESC))
        ws.row_dimensions[r].height = 20
    totais = [("TOTAL", None), (f"=SUM({V}{r0 + 1}:{V}{rt - 1})", BRL0), (f"=SUM({PA}{r0 + 1}:{PA}{rt - 1})", PCT),
              (f"=SUM({PI}{r0 + 1}:{PI}{rt - 1})", PCT), (None, None), (f"=SUM({AP}{r0 + 1}:{AP}{rt - 1})", BRL0)]
    for (a, b), (x, fmt) in zip(spans, totais):
        ws.merge_cells(start_row=rt, start_column=a, end_row=rt, end_column=b)
        cc = ws.cell(rt, a, x)
        cc.font = Font(name=F, size=9, bold=True, color=AMARELO)
        cc.alignment = Alignment(horizontal="left" if a == X0 else "right", indent=1)
        if fmt:
            cc.number_format = fmt
        for cx in range(a, b + 1):
            ws.cell(rt, cx).fill = PatternFill("solid", fgColor=BORDA_ESC)
    ws.row_dimensions[rt].height = 20
    sinal(ws, f"{DV}{r0 + 1}:{DV}{rt - 1}")
    ws.print_area = f"A1:{L(ULT + 1)}{rt + 1}"
    return r0, rt
