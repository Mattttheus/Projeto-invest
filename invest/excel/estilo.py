"""Identidade visual única da planilha: cores, fontes, formatos numéricos e blocos reutilizáveis."""
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo

# paleta de sistema SaaS/ERP: grafite (estrutura), azul (primária), verde/âmbar/vermelho (status)
F = "Segoe UI"
SIDEBAR, NAVY, AZUL = "0F172A", "1E293B", "2563EB"
TEAL, AMBAR, VERMELHO, VERDE = "0D9488", "D97706", "DC2626", "16A34A"
TEXTO, CINZA_TXT, BORDA_HEX, FUNDO_HEX = "0F172A", "64748B", "E2E8F0", "F1F5F9"
# paleta executiva enxuta: azul = dado principal, cinza = referência/comparação
CINZA_SERIE = "CBD5E1"
CORES = [AZUL, CINZA_SERIE, "1E3A8A", "94A3B8"]

f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_norm = Font(name=F, size=10)
f_input = Font(name=F, size=10, color="0000FF")      # entrada (vem de dados/)
f_link = Font(name=F, size=10, color="008000")       # vínculo com outra aba
f_bold = Font(name=F, size=10, bold=True)
f_note = Font(name=F, size=9, color="444444")
f_sub = Font(name=F, size=9, italic=True, color="666666")

fill_head = PatternFill("solid", fgColor=NAVY)
fill_head2 = PatternFill("solid", fgColor=NAVY)
fill_input = PatternFill("solid", fgColor="FFF7D6")
fill_total = PatternFill("solid", fgColor="E8EDF4")
fill_zebra = PatternFill("solid", fgColor="F8FAFC")
FUNDO = PatternFill("solid", fgColor=FUNDO_HEX)
BRANCO = PatternFill("solid", fgColor="FFFFFF")

thin = Side(style="thin", color="D8DFE9")
b_all = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

BRL = '"R$" #,##0.00;[Red]-"R$" #,##0.00;"-"'
BRL0 = '"R$" #,##0;[Red]-"R$" #,##0;"-"'
BRL4 = '"R$" #,##0.0000'
PCT = '0.0%;[Red]-0.0%;"-"'
QTD = '#,##0;[Red]-#,##0;"-"'
DATE = "dd/mm/yyyy"
MESES = '0" meses";;"atingida"'

# abas na ordem de leitura: resultados primeiro, entradas depois
ABAS = ["Painel", "Carteira", "Análises", "Cadastro", "Ativos", "Lançamentos", "Proventos", "Premissas", "Verificar"]


def ref(aba):
    return f"'{aba}'" if not aba.isalnum() else aba


NAVBAR_COLUNAS = 40            # a faixa escura da linha 1 cobre a largura do menu, mesmo em abas estreitas
MODO = {"google": False}      # True ao gerar a versão para Google Planilhas (sem VBA/formas)
MENU_GOOGLE = ["Painel", "Carteira", "Análises", "✚ Lançar", "Lançamentos", "Proventos", "Premissas", "Verificar"]


def navbar_texto(ws, ultima_col):
    """Google Planilhas: menu como texto na linha 1 (o Apps Script transforma cada item em link ao abrir)."""
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=ultima_col)
    c = ws.cell(1, 2, "◆ InvestERP        " + "     ".join(MENU_GOOGLE))
    c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
    c.alignment = Alignment(vertical="center")


def faixa(ws, titulo, subtitulo, ncols, cor=None):
    """Cabeçalho de sistema: linha 1 = barra de navegação escura (itens inseridos em navegacao.py),
    linha 2 = título, linha 3 = subtítulo. Dados começam na linha 4; linhas 1-3 ficam fixas ao rolar."""
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    escuro, claro = PatternFill("solid", fgColor=SIDEBAR), PatternFill("solid", fgColor="FFFFFF")
    for c in range(1, max(ncols, NAVBAR_COLUNAS) + 1):
        ws.cell(1, c).fill = escuro
    for c in range(1, ncols + 1):
        ws.cell(2, c).fill = claro
        ws.cell(2, c).border = Border(bottom=Side(style="medium", color=AZUL))
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 34
    ws.row_dimensions[3].height = 16
    if MODO["google"]:
        navbar_texto(ws, max(ncols, 26))
    else:
        marca = ws.cell(1, 1, "  ◆ InvestERP")
        marca.font = Font(name=F, size=12, bold=True, color="FFFFFF")
        marca.alignment = Alignment(vertical="center")
    t = ws.cell(2, 1, titulo)
    t.font = Font(name=F, size=15, bold=True, color=TEXTO)
    t.alignment = Alignment(vertical="center", indent=1)
    s = ws.cell(3, 1, subtitulo)
    s.font = Font(name=F, size=8, italic=True, color=CINZA_TXT)
    s.alignment = Alignment(vertical="top", indent=1)


def cabecalho(ws, row, cols, fill=fill_head):
    for i, h in enumerate(cols, 1):
        c = ws.cell(row, i, h)
        c.font, c.fill, c.alignment, c.border = f_head, fill, center, b_all
    ws.row_dimensions[row].height = 32


def tabela(ws, nome, primeira_linha, ultima_linha, ncols):
    """Tabela do Excel (filtros e ordenação) sem sobrescrever a formatação própria."""
    if ultima_linha <= primeira_linha:
        return
    t = Table(displayName=nome, ref=f"A{primeira_linha}:{L(ncols)}{ultima_linha}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)


def larguras(ws, ws_w):
    for i, w in enumerate(ws_w, 1):
        ws.column_dimensions[L(i)].width = w


def cel(ws, r, col, v, fmt=None, font=f_norm, fill=None):
    c = ws.cell(r, col, v)
    c.font, c.border = font, b_all
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def linha_total(ws, r, ncols, rotulo="TOTAL"):
    for col in range(1, ncols + 1):
        c = ws.cell(r, col)
        c.font, c.fill, c.border = f_bold, fill_total, b_all
    ws.cell(r, 1, rotulo)


def notas(ws, row, itens, titulo="Como ler"):
    ws.cell(row, 1, titulo).font = f_bold
    for k, n in enumerate(itens):
        ws.cell(row + 1 + k, 1, "• " + n).font = f_note


def nulo(v):
    return None if pd.isna(v) else v
